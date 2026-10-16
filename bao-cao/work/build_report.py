from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path.cwd(); WORK=ROOT/'bao-cao/work'; IMG=WORK/'images'
OUT=ROOT/'bao-cao/Bao_cao_He_thong_IoT_Nha_thong_minh.docx'
E=json.loads((WORK/'evidence.json').read_text(encoding='utf8'))
R=json.loads((WORK/'runtime-check.json').read_text(encoding='utf8'))
D=Document()
members=[('Nguyễn Thùy Linh','23810310378'),('Lê Hải Nam','23810310377'),('Nguyễn Xuân Đạt','23810310361'),('Vũ Tiến Mạnh','23810310375'),('Đinh Việt Dũng','23810310370'),('Nguyễn Đắc Tuấn Nghĩa','23810310343')]
figures=[]; tables=[]; bookmark_id=0

def setup_section(s):
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.3);s.left_margin=Cm(3);s.right_margin=Cm(2)
    s.header_distance=Cm(1.2);s.footer_distance=Cm(1.1)
setup_section(D.sections[0])
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2']:
    if name not in D.styles:continue
    s=D.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
n=D.styles['Normal'];n.font.size=Pt(13)
n.paragraph_format.line_spacing=1.3;n.paragraph_format.space_after=Pt(6)
n.paragraph_format.first_line_indent=Cm(.8);n.paragraph_format.widow_control=True
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
    s=D.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.first_line_indent=Cm(0)
    s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(8);s.paragraph_format.keep_with_next=True
    s.paragraph_format.line_spacing=1.15
D.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.styles['Caption'].font.size=Pt(12);D.styles['Caption'].font.italic=True
D.styles['Caption'].paragraph_format.first_line_indent=Cm(0)
D.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.styles['Caption'].paragraph_format.line_spacing=1.15
for name in ['TOC 1','TOC 2']:
    if name in D.styles:
        s=D.styles[name];s.font.size=Pt(12);s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.1

# Remove template title rules and theme-font overrides before Word refreshes fields.
for style in D.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    if style.type == 1:
        rf=style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+key),None)
        for key in ['ascii','hAnsi','eastAsia','cs']:
            rf.set(qn('w:'+key),'Times New Roman')

def field(p,code,text='1'):
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    r=p.add_run();a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=' '+code+' ';r._r.append(a)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a)
    p.add_run(text)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)
def bookmark(p,prefix):
    global bookmark_id
    bookmark_id+=1;name=f'{prefix}_{bookmark_id}'
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmark_id));a.set(qn('w:name'),name);p._p.insert(0,a)
    z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(bookmark_id));p._p.append(z);return name
def P(*paras):
    for text in paras:
        p=D.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
def H(text,level=2):return D.add_paragraph(text,style=f'Heading {level}')
def chapter(text):
    p=H(text,1);p.paragraph_format.page_break_before=True;return p
def centered(text,size=14,bold=False,after=8,before=0):
    p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(before)
    r=p.add_run(text);r.bold=bold;r.font.size=Pt(size);return p
def front_title(text):
    p=centered(text,16,True,after=16);p.paragraph_format.keep_with_next=True;return p
def T(caption,headers,rows,widths=None):
    if caption:
        p=D.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True
        b=bookmark(p,'tbl');tables.append((caption,b))
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if not widths:widths=[16/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();a=OxmlElement('w:cantSplit');pr.append(a)
        if ri==0:a=OxmlElement('w:tblHeader');pr.append(a)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'5');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for e,v in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                x=OxmlElement('w:'+e);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');margins.append(x)
            cp.append(margins)
            if ri==0:
                a=OxmlElement('w:shd');a.set(qn('w:fill'),'E9EFF5');cp.append(a)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.12
                if ri < 3 or ri == len(t.rows)-2:
                    p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if widths[ci]<2.8 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.size=Pt(11.5);r.bold=ri==0
    D.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def F(file,caption,width=15.5,maxheight=18):
    source=Path(file)
    if not source.is_absolute():source=IMG/source
    im=Image.open(source);w=min(width,maxheight*im.width/im.height)
    p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
    r=p.add_run();r.add_picture(str(source),width=Cm(w))
    dp=r._r.xpath('.//wp:docPr')
    if dp:dp[0].set('descr',caption)
    p=D.add_paragraph(caption,style='Caption');figures.append((caption,bookmark(p,'fig')))
def code(text):
    for line in text.strip('\n').splitlines():
        p=D.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02
        r=p.add_run(line or ' ');r.font.name='Consolas';r.font.size=Pt(9.5)
def page_numbers(section,fmt='decimal',start=1):
    section.footer.is_linked_to_previous=False
    a=OxmlElement('w:pgNumType');a.set(qn('w:fmt'),fmt);a.set(qn('w:start'),str(start));section._sectPr.append(a)
    p=section.footer.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(p,'PAGE')
    for r in p.runs:r.font.size=Pt(11)

# Cover, using the academic page format of the supplied course reports.
centered('BỘ CÔNG THƯƠNG',14,True,after=2)
centered('TRƯỜNG ĐẠI HỌC ĐIỆN LỰC',15,True,after=2)
centered('KHOA CÔNG NGHỆ THÔNG TIN',14,True,after=36)
centered('BÁO CÁO CHUYÊN ĐỀ HỌC PHẦN',17,True,after=6)
centered('HỆ THỐNG IoT VÀ ỨNG DỤNG',16,True,after=28)
title=centered('THIẾT KẾ HỆ THỐNG GIÁM SÁT\nVÀ ĐIỀU KHIỂN NHÀ THÔNG MINH\nTÍCH HỢP TRỢ LÝ AI',19,True,after=28)
title.style=D.styles['Title'];title.paragraph_format.first_line_indent=Cm(0)
for r in title.runs:r.font.name='Times New Roman';r.font.size=Pt(19);r.font.color.rgb=RGBColor(0,0,0);r.bold=True
centered('Giảng viên hướng dẫn  ThS. Phạm Văn Hiệp',14,False,after=18)
centered('Sinh viên thực hiện',14,True,after=8)
T(None,['Họ và tên','Mã sinh viên','Lớp'],[(n,m,'D18CNPM5') for n,m in members],[7.2,4.5,4.3])
centered('Hà Nội, tháng 10 năm 2026',13,False,before=20)

s=D.add_section(WD_SECTION_START.NEW_PAGE);setup_section(s);page_numbers(s,'lowerRoman',1)
front_title('LỜI CẢM ƠN')
P('Chúng em xin trân trọng cảm ơn ThS. Phạm Văn Hiệp đã giảng dạy và hướng dẫn chúng em trong học phần Hệ thống IoT và ứng dụng. Những kiến thức về cảm biến, kết nối thiết bị, giao thức truyền thông và xây dựng ứng dụng đã giúp nhóm hình thành và triển khai hệ thống nhà thông minh trong báo cáo này.',
'Chúng em cảm ơn các thầy cô Khoa Công nghệ Thông tin, Trường Đại học Điện lực đã trang bị nền tảng kiến thức lập trình, cơ sở dữ liệu và mạng máy tính. Trong quá trình thực hiện, các thành viên trong nhóm cùng trao đổi, lắp ráp mô hình, kiểm tra chức năng và hoàn thiện phần mềm để kết nối các nội dung đã học với một bài toán cụ thể.',
'Báo cáo trình bày quá trình thiết kế, triển khai và đánh giá hệ thống giám sát phòng khách sử dụng ESP32, MQTT, Node.js, MySQL và giao diện web, đồng thời tích hợp trợ lý AI chạy cục bộ. Nhóm mong nhận được ý kiến góp ý của giảng viên để cải thiện độ ổn định, tính dễ sử dụng và khả năng mở rộng của hệ thống.')
centered('Nhóm sinh viên lớp D18CNPM5',13,True,before=18)
D.add_page_break();front_title('TÓM TẮT ĐỀ TÀI')
P('Đề tài xây dựng một mô hình nhà thông minh ở phạm vi phòng khách, trong đó ESP32 thu nhận nhiệt độ, độ ẩm, chuyển động và trạng thái ánh sáng; điều khiển relay, đèn LED và còi; đồng thời trao đổi dữ liệu với máy chủ qua WiFi và MQTT. Hệ thống kết hợp tự động hóa tại thiết bị với giám sát, lưu trữ và tương tác trên trình duyệt.',
'Phần máy chủ sử dụng Node.js và Express để cung cấp REST API, MySQL để lưu lịch sử và thói quen người dùng, WebSocket để cập nhật dữ liệu cho dashboard React. Trợ lý AI có hai nhánh xử lý: nhận diện lệnh đèn và còi bằng quy tắc; trả lời câu hỏi về ngôi nhà bằng mô hình llama3.1 chạy trên Ollama với ngữ cảnh lấy từ cơ sở dữ liệu.',
'Dữ liệu lịch sử của hai phiên chạy ngày 01/09/2026 và 08/09/2026 gồm 2.077 bản ghi MQTT, trong đó có 1.950 bản ghi trạng thái và 127 bản ghi lệnh. Có 304 bản ghi nhiệt độ trong khoảng 29,4–35,0°C và 303 bản ghi độ ẩm trong khoảng 54–88%. Các con số này mô tả tập dữ liệu đã lưu, không thay thế phép hiệu chuẩn cảm biến hay phép đo tỷ lệ tiết kiệm điện.',
'Các thử nghiệm chức năng trước đây được nhóm xác nhận thành công, bao gồm giám sát cảm biến, điều khiển tự động và thủ công, lưu lịch sử, cập nhật giao diện và hỏi đáp AI. Lần kiểm tra ngày 15/10/2026 tiếp tục xác nhận kết nối MQTT, REST API, WebSocket, quá trình build giao diện và phản hồi chatbot. Báo cáo đồng thời phân tích giới hạn của kiến trúc hiện tại và đề xuất hướng hoàn thiện.')
P('Từ khóa: nhà thông minh, ESP32, ESPHome, MQTT, Node.js, React, MySQL, Ollama.')
D.add_page_break();front_title('MỤC LỤC')
p=D.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);field(p,'TOC \\o "1-2" \\h \\z \\u','Mục lục')
D.add_page_break();front_title('DANH MỤC TỪ VIẾT TẮT')
T(None,['Từ viết tắt','Tên gọi','Ý nghĩa trong đề tài'],[
('IoT','Internet of Things','Kết nối thiết bị vật lý với ứng dụng'),('MCU','Microcontroller Unit','Vi điều khiển ESP32'),('GPIO','General Purpose Input Output','Chân tín hiệu vào và ra'),('PIR','Passive Infrared','Cảm biến phát hiện chuyển động'),('LDR','Light Dependent Resistor','Phần tử quang trở trong module ánh sáng'),('MQTT','MQTT','Giao thức truyền bản tin publish/subscribe'),('LAN','Local Area Network','Mạng nội bộ'),('TCP','Transmission Control Protocol','Giao thức vận chuyển'),('HTTP','Hypertext Transfer Protocol','Giao tiếp yêu cầu và phản hồi'),('REST','Representational State Transfer','Kiểu tổ chức API'),('API','Application Programming Interface','Giao diện lập trình ứng dụng'),('WS','WebSocket','Kết nối đẩy dữ liệu tới dashboard'),('JSON','JavaScript Object Notation','Dạng biểu diễn dữ liệu'),('OTA','Over The Air','Cập nhật firmware qua mạng'),('LLM','Large Language Model','Mô hình ngôn ngữ dùng cho hỏi đáp'),('AI','Artificial Intelligence','Trí tuệ nhân tạo')],[2.4,5.4,8.2])
D.add_page_break();front_title('DANH MỤC HÌNH ẢNH')
figure_list_anchor=D.add_paragraph('FIGURE_LIST_ANCHOR')
D.add_page_break();front_title('DANH MỤC BẢNG BIỂU')
table_list_anchor=D.add_paragraph('TABLE_LIST_ANCHOR')

s=D.add_section(WD_SECTION_START.NEW_PAGE);setup_section(s);page_numbers(s,'decimal',1)
H('CHƯƠNG 1 TỔNG QUAN VỀ ĐỀ TÀI',1)
H('Lý do chọn đề tài')
P('Trong một phòng sinh hoạt, người sử dụng thường cần biết nhiệt độ, độ ẩm và tình trạng thiết bị, đồng thời muốn bật tắt các cơ cấu chấp hành mà không phải thao tác trực tiếp tại từng công tắc. Nếu dữ liệu chỉ được quan sát tại chỗ, việc theo dõi diễn biến theo thời gian và đối chiếu các sự kiện trở nên khó khăn. Đây là bài toán phù hợp để vận dụng kiến thức IoT từ thu thập tín hiệu đến truyền dữ liệu, lưu trữ và xây dựng giao diện.',
'Nhóm lựa chọn mô hình phòng khách vì phạm vi này đủ rõ để xây dựng và kiểm thử một chuỗi xử lý hoàn chỉnh. DHT11 cung cấp hai đại lượng môi trường; PIR tạo sự kiện chuyển động; module LDR cung cấp tín hiệu ánh sáng dạng nhị phân. Relay, LED và còi giúp thể hiện kết quả điều khiển trên mô hình thật, qua đó kiểm tra được quan hệ giữa cảm biến, chương trình nhúng và ứng dụng giám sát.',
'Bên cạnh các thao tác truyền thống, nhóm bổ sung trợ lý AI để người dùng hỏi về dữ liệu đang có và gửi một số lệnh ngắn. Phần mở rộng này tạo cơ hội nghiên cứu cách cung cấp ngữ cảnh thực tế cho mô hình ngôn ngữ, đồng thời giữ những lệnh điều khiển đơn giản trong nhánh xử lý xác định. Nhờ đó, đề tài bao quát được cả phần cứng, mạng, phần mềm và tương tác người dùng.')
H('Mục tiêu của hệ thống')
P('Mục tiêu tổng quát là xây dựng hệ thống giám sát và điều khiển nhà thông minh hoạt động trong mạng nội bộ, cho phép theo dõi trạng thái phòng khách qua trình duyệt và duy trì các phản ứng tự động tại ESP32. Máy chủ tập trung vào lưu trữ, cung cấp dữ liệu, tạo gợi ý và hỗ trợ hỏi đáp.',
'Về thu thập dữ liệu, hệ thống phải đọc được nhiệt độ, độ ẩm, chuyển động và ánh sáng, tổ chức bản tin theo từng topic MQTT và chuyển được dữ liệu về cơ sở dữ liệu. Về điều khiển, hệ thống phải thực hiện được logic tự động theo PIR và ánh sáng, tiếp nhận lệnh bật tắt relay, LED, còi và chuyển đổi chế độ thủ công.',
'Về ứng dụng, dashboard phải hiển thị các giá trị gần nhất, biểu đồ dữ liệu, trạng thái kết nối và công cụ điều khiển. Hệ thống cần quản lý các ghi chú thói quen để làm ngữ cảnh tham chiếu và cung cấp chatbot hỏi đáp dựa trên dữ liệu cảm biến, gợi ý, bất thường và thói quen đã lưu.')
H('Đối tượng và phạm vi thực hiện')
P('Đối tượng nghiên cứu là chuỗi kết nối giữa một nút ESP32, các cảm biến và cơ cấu chấp hành trong mô hình phòng khách, dịch vụ MQTT, máy chủ ứng dụng và trình duyệt. Mô hình có một nút điều khiển trung tâm ở phía thiết bị; các dịch vụ xử lý và lưu trữ được triển khai trên máy tính trong mạng LAN.',
'Phạm vi điều khiển trực tiếp gồm một kênh relay ở GPIO26, LED ở GPIO25 và còi ở GPIO33. Mặc dù phần cứng dùng module relay hai kênh, cấu hình hiện tại chỉ khai báo một kênh relay. Cảm biến ánh sáng được đọc qua ngõ ra số ở GPIO34; hệ thống chưa đo cường độ ánh sáng theo đơn vị lux.',
'Đề tài tập trung vào mô hình học tập và kiểm chứng các luồng chức năng. Việc mở rộng nhiều phòng, xây dựng tài khoản người dùng, triển khai truy cập từ Internet, hiệu chuẩn cảm biến với thiết bị chuẩn và đo công suất tiêu thụ là các hướng phát triển tiếp theo. Các thao tác ghi thói quen hiện tại chưa tạo ra lịch tự động thực thi.')
H('Phương pháp thực hiện')
P('Nhóm tiếp cận bài toán theo từng khối chức năng: lựa chọn cảm biến và chân kết nối; xây dựng cấu hình ESPHome; tổ chức truyền thông MQTT; triển khai máy chủ, cơ sở dữ liệu và dashboard; sau đó tích hợp chatbot. Cách làm này giúp kiểm tra riêng từng khối trước khi ghép thành hệ thống.',
'Việc đánh giá kết hợp ba nguồn: cấu hình và mã nguồn triển khai; bản ghi dữ liệu từ các phiên chạy đã có; thử nghiệm lại các dịch vụ phần mềm. Trạng thái phần cứng được đối chiếu với topic phản hồi, thay vì chỉ suy ra từ việc API đã nhận một yêu cầu điều khiển. Các phép kiểm tra logic dùng các giá trị ở hai phía của ngưỡng để tránh bỏ sót sai lệch tại biên.',
'Kết quả định lượng được thống kê từ dữ liệu MySQL. Khi phân tích, nhóm tách bản tin trạng thái khỏi bản tin lệnh và tách dữ liệu lịch sử khỏi dữ liệu retained nhận lại sau khi máy chủ kết nối. Cách phân loại này giúp diễn giải đúng số bản ghi và tránh coi các bản tin phát lại là phép đo mới.')
H('Cấu trúc báo cáo')
P('Báo cáo gồm bốn chương: tổng quan đề tài; cơ sở lý thuyết và công nghệ; phân tích thiết kế và triển khai; kết quả thử nghiệm và đánh giá. Phần cuối gồm kết luận, tài liệu tham khảo và phụ lục cấu hình, vận hành.')

chapter('CHƯƠNG 2 CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ')
H('Mô hình chức năng của hệ thống IoT')
P('Trong đề tài, hệ thống được mô tả theo bốn nhóm chức năng: cảm nhận và chấp hành; kết nối mạng; xử lý và lưu trữ; ứng dụng. Đây là cách phân chia để giải thích vai trò của các khối trong mô hình, không phải phép đồng nhất với bảy lớp OSI. TCP được xem là giao thức vận chuyển dùng bên dưới các giao thức ứng dụng.',
'Ở phía thiết bị, ESP32 đọc tín hiệu vật lý và thực hiện logic điều khiển tại chỗ. Phía mạng dùng WiFi và mạng IP để đưa bản tin đến broker. Phía xử lý gồm broker, backend, MySQL và Ollama. Phía ứng dụng là dashboard trên trình duyệt. Một trạng thái cảm biến có thể đi xuyên qua cả bốn nhóm, trong khi thao tác tự động tại ESP32 có thể diễn ra mà không chờ xử lý từ máy chủ.')
T('Bảng 2.1 Các khối chức năng trong mô hình',['Khối','Thành phần','Vai trò'],[
('Cảm nhận và chấp hành','ESP32, DHT11, PIR, LDR, relay, LED, còi','Thu tín hiệu và tác động tại mô hình'),('Kết nối','WiFi, IP, TCP, MQTT','Truyền trạng thái và lệnh'),('Xử lý và lưu trữ','Mosquitto, Node.js, MySQL, Ollama','Tiếp nhận, lưu lịch sử, xử lý yêu cầu'),('Ứng dụng','React, REST, WebSocket','Hiển thị và tương tác với người dùng')],[3.4,5.6,7])
H('ESP32 và nền tảng ESPHome')
P('ESP32 là vi điều khiển trung tâm của mô hình. Các chân GPIO cho phép tiếp nhận tín hiệu số hoặc xuất mức điều khiển đến mạch ngoài. Việc chọn chân phải xét chức năng của từng nhóm chân; riêng GPIO34 thuộc nhóm chân chỉ có khả năng nhập trên ESP32 và không có điện trở kéo lên, kéo xuống nội tích hợp [1]. Trong cấu hình của đề tài, chân này được dùng để đọc ngõ ra module ánh sáng.',
'ESPHome cho phép khai báo thiết bị, cảm biến, công tắc, kết nối mạng và hành động tự động bằng YAML. Cấu hình nhóm sử dụng biến thể esp32 và framework esp-idf, khai báo DHT11 cùng các binary_sensor, switch. Topic trạng thái và topic lệnh được cấu hình rõ ràng để phần mềm máy chủ có thể dùng một giao ước thống nhất [2].',
'Một ưu điểm của cách triển khai này là logic bật LED và còi nằm trực tiếp trên ESP32. Khi hệ thống đã khởi động và vẫn được cấp nguồn, các nhánh xử lý PIR không cần gửi dữ liệu lên máy chủ rồi chờ lệnh quay về. Tuy vậy, giám sát trên web và truy vấn lịch sử vẫn phụ thuộc vào các dịch vụ máy chủ và kết nối mạng.')
H('Các cảm biến và cơ cấu chấp hành')
P('DHT11 được dùng để thu nhiệt độ và độ ẩm tương đối. Trong ESPHome, một thành phần dht có thể tạo ra hai kênh sensor với topic trạng thái riêng. Nhóm chỉ định model DHT11 và chân dữ liệu GPIO4. Do cấu hình không khai báo update_interval, chu kỳ mặc định theo tài liệu thành phần DHT là 60 giây [3]. Chu kỳ thực tế trong lịch sử có thể bị ảnh hưởng bởi gián đoạn chạy thử hoặc thời điểm kết nối lại.',
'PIR HC-SR501 đóng vai trò nguồn sự kiện chuyển động, còn module LDR HW-072 đưa ra trạng thái nhị phân theo ngưỡng của mạch so sánh. Chương trình đọc hai ngõ ra này như binary_sensor. Vì dữ liệu ánh sáng chỉ là ON hoặc OFF, nhóm sử dụng nó để quyết định một điều kiện logic; không diễn giải giá trị này thành một phép đo độ rọi liên tục.',
'Relay tách tín hiệu điều khiển GPIO khỏi mạch tải qua cơ cấu đóng cắt của module. Trong cấu hình, relay ở GPIO26 được đặt inverted: true để ánh xạ lệnh bật sang mức điện phù hợp với module. LED và còi được khai báo dưới dạng switch GPIO, lần lượt tại GPIO25 và GPIO33. Trạng thái phần mềm phản ánh lệnh xuất GPIO; hệ thống chưa có cảm biến dòng hoặc tiếp điểm phản hồi để chứng minh tải đã hoạt động vật lý.')
F('components.png','Hình 2.1 Các thành phần phần cứng sử dụng trong mô hình',15.5,13)
H('WiFi và giao tiếp trong mạng nội bộ')
P('WiFi được chọn vì mô hình triển khai trong phạm vi một phòng và có thể dùng bộ phát mạng sẵn có. ESP32 và máy chủ phải truy cập được broker theo địa chỉ IP trong LAN. Trong cấu hình nhóm cung cấp, broker có địa chỉ 192.168.0.102, cổng 1883; phía backend dùng localhost vì kết nối đến broker được ánh xạ trên cùng máy chủ.',
'Cần phân biệt địa chỉ dùng trên máy chủ với địa chỉ dùng trên thiết bị khác. localhost trong trình duyệt luôn chỉ đến chính thiết bị mở trình duyệt. Vì dashboard hiện đặt API_BASE là http://localhost:3000 và WS_URL là ws://localhost:8081, việc chạy trên máy chủ là thuận tiện; khi mở từ điện thoại hoặc máy khác, các giá trị này cần trỏ về địa chỉ LAN của máy chủ.',
'Kiến trúc cục bộ giúp các dịch vụ chính không phải gửi dữ liệu ra một nền tảng đám mây bên ngoài. Tuy nhiên, khả năng hoạt động vẫn phụ thuộc nguồn cấp ESP32, bộ phát WiFi, broker và máy tính chạy ứng dụng. Việc truy cập từ xa cần được thiết kế như một chức năng riêng có kiểm soát truy cập phù hợp.')
H('Giao thức MQTT')
P('MQTT dùng mô hình publish/subscribe: bên gửi công bố bản tin lên topic và broker chuyển bản tin đến các bên đã đăng ký. Cách giao tiếp này phù hợp với các thiết bị có tài nguyên hạn chế và các thông điệp ngắn. MQTT có các mức QoS để lựa chọn cơ chế chuyển giao bản tin [4]. Trong mã nguồn backend của đề tài, subscribe và publish không truyền tùy chọn QoS riêng.',
'Cây topic của hệ thống bắt đầu bằng nha/phongkhach. Các topic kết thúc bằng /state biểu diễn trạng thái, còn /set biểu diễn lệnh. Backend đăng ký nha/phongkhach/# nên nhận cả hai loại. MQTT hỗ trợ bản tin retained để phía đăng ký mới nhận lại trạng thái đã lưu; tài liệu ESPHome cũng cho phép tùy chỉnh state_topic, command_topic và hành vi retain [2].',
'Nhóm dùng retained để duy trì giá trị tham khảo khi khởi động lại phần mềm. Tuy nhiên, retained không chứng minh thiết bị đang trực tuyến và không cung cấp thời điểm đo gốc trong payload hiện tại. Vì vậy, phần lưu trữ và giao diện cần phân biệt thời điểm nhận lại bản tin với thời điểm cảm biến thực sự tạo phép đo.')
H('REST API và WebSocket')
P('REST API được tổ chức theo các đường dẫn để lấy lịch sử, gửi lệnh, quản lý thói quen và gọi chatbot. Express định tuyến yêu cầu dựa trên phương thức HTTP và đường dẫn; các tham số trên đường dẫn có thể được sử dụng để xác định loại dữ liệu cần truy vấn [5]. Trong đề tài, GET /api/lichsu/:loai ánh xạ tên loại sang topic cảm biến tương ứng.',
'WebSocket được dùng cho các cập nhật đẩy từ backend tới dashboard. Khác với việc tải lại toàn bộ trang, dashboard giữ một kết nối và nhận các JSON mà máy chủ phát ra. Mỗi bản tin MQTT được ghi MySQL thành công sẽ được broadcast qua cổng 8081. Việc tách REST và WebSocket giúp sử dụng REST để khởi tạo dữ liệu, còn WebSocket để cập nhật khi có bản tin mới.',
'Trạng thái kết nối WebSocket chỉ cho biết trình duyệt đã kết nối đến backend. Để đánh giá toàn bộ đường truyền, còn cần biết backend có kết nối broker và ESP32 có hoạt động hay không. Đây là ba lớp trạng thái khác nhau; chỉ hiển thị một nhãn kết nối chung có thể khiến người dùng hiểu sai độ mới của dữ liệu.')
H('Node.js và cơ sở dữ liệu MySQL')
P('Backend viết bằng JavaScript, chạy trên Node.js và kết hợp các thư viện mqtt, express, mysql2/promise, ws, cors và dotenv. Chương trình nhận bản tin, ghi dữ liệu, phát thông báo và xử lý HTTP trong một dịch vụ. Các thao tác truy vấn cơ sở dữ liệu dùng lời gọi bất đồng bộ để chờ kết quả trước khi tiếp tục xử lý.',
'MySQL lưu trữ lịch sử dưới dạng các hàng có topic, giá trị và thời gian nhận. Cách thiết kế chung một bảng cho nhiều loại bản tin giúp thêm topic mà không phải thêm cột. Đổi lại, trường giá trị là chuỗi nên các truy vấn thống kê nhiệt độ và độ ẩm phải chuyển kiểu số; đồng thời cần chuẩn hóa payload trước khi sử dụng cho các phép tính.',
'Một bảng riêng lưu danh mục và nội dung thói quen người dùng. Backend đưa dữ liệu này vào ngữ cảnh chatbot. Trong phiên bản hiện tại, đây là thông tin do người dùng ghi lại; hệ thống chưa tự học thói quen từ dữ liệu và chưa có bộ lập lịch biến một câu mô tả thói quen thành lệnh điều khiển tự động.')
H('Giao diện React và trợ lý AI cục bộ')
P('React tổ chức giao diện thành các thành phần như thẻ cảm biến, khối điều khiển relay, biểu đồ, danh mục thiết bị và trang hội thoại. State lưu các giá trị được hiển thị; các effect thực hiện tác vụ tải dữ liệu hoặc thiết lập kết nối khi thành phần được gắn vào trang [6]. Dashboard của đề tài dùng cách tổ chức này để kết hợp dữ liệu REST ban đầu với dữ liệu WebSocket.',
'Ollama cung cấp API HTTP để gọi mô hình ngôn ngữ chạy cục bộ. API /api/chat nhận model và messages; tùy chọn stream: false yêu cầu trả về một phản hồi hoàn chỉnh thay vì chuỗi kết quả theo thời gian [7]. Backend đang gọi model llama3.1:latest tại cổng 11434 và hướng dẫn mô hình trả lời ngắn bằng tiếng Anh.',
'Phần điều khiển qua câu lệnh và phần hỏi đáp có vai trò khác nhau. Câu lệnh đèn, còi được nhận diện bằng regex trước khi gọi mô hình; các câu hỏi khác được gửi kèm dữ liệu tham chiếu đến LLM. Kiến trúc hiện tại không huấn luyện lại mô hình, không tự học từ lịch sử hội thoại và không để LLM tự sinh lệnh MQTT để thực thi.')

chapter('CHƯƠNG 3 PHÂN TÍCH THIẾT KẾ VÀ TRIỂN KHAI')
H('Yêu cầu chức năng')
P('Hệ thống được thiết kế cho người sử dụng trong mạng nội bộ. Người dùng cần quan sát thông tin phòng khách, nhận gợi ý từ dữ liệu, điều khiển một số đầu ra và hỏi trợ lý. Các yêu cầu được chia theo luồng thu thập, điều khiển và khai thác dữ liệu để đối chiếu với các thành phần triển khai.')
T('Bảng 3.1 Các yêu cầu chức năng',['Mã','Chức năng','Cách thực hiện'],[
('CN01','Giám sát nhiệt độ và độ ẩm','Đọc DHT11, gửi hai topic, hiển thị thẻ và biểu đồ'),('CN02','Giám sát chuyển động và ánh sáng','Đọc PIR, LDR dưới dạng trạng thái số'),('CN03','Tự động điều khiển tại thiết bị','Xử lý on_press và on_release của PIR trong ESPHome'),('CN04','Điều khiển relay từ web','POST /api/relay rồi publish relay/set'),('CN05','Điều khiển LED và còi bằng câu lệnh','Nhận diện câu lệnh, bật chế độ thủ công và publish lệnh'),('CN06','Theo dõi chế độ vận hành','Nhận thucong/state; cho phép quay lại tự động'),('CN07','Lưu và đọc lịch sử','MySQL và GET /api/lichsu/:loai'),('CN08','Gợi ý và phát hiện bất thường','Ngưỡng cố định và truy vấn dữ liệu lịch sử'),('CN09','Quản lý thói quen','Thêm, xem, xóa các nội dung theo danh mục'),('CN10','Hỏi đáp có ngữ cảnh','Đọc dữ liệu, tạo context, gọi Ollama')],[1.5,5.1,9.4])
H('Yêu cầu vận hành')
P('Hệ thống cần duy trì đường đi của dữ liệu từ cảm biến đến cơ sở dữ liệu, đồng thời thể hiện được trạng thái đầu ra sau thao tác điều khiển. Dashboard phải có cách xử lý mất kết nối và không làm mất dữ liệu lịch sử khi đóng trình duyệt. Trong mã hiện tại, WebSocket được kết nối lại sau 3 giây nếu đường truyền bị đóng.',
'Máy chủ cần chạy đồng thời broker, MySQL, backend và dashboard; Ollama cần sẵn sàng khi sử dụng nhánh hỏi đáp AI. Cổng 1883 dùng cho MQTT, 3000 cho REST API, 8081 cho WebSocket và 11434 cho Ollama. Cổng 5173 được sử dụng khi chạy Vite phục vụ phát triển giao diện.',
'Cấu hình xác thực được tách khỏi nội dung công khai của báo cáo. Mật khẩu cơ sở dữ liệu nằm trong biến môi trường của backend; các mật khẩu WiFi, MQTT, OTA và khóa API của ESPHome được biểu diễn bằng tham chiếu secrets trong phụ lục. Phụ lục giữ nguyên logic điều khiển để có thể đối chiếu với phiên bản của nhóm.')
H('Kiến trúc tổng thể')
P('Kiến trúc có hai điểm xử lý chính. ESP32 xử lý các phản ứng theo sự kiện ngay tại mô hình; backend xử lý lưu trữ, giao diện và trợ lý. MQTT broker là điểm trung chuyển hai chiều. Dashboard không trực tiếp truy cập GPIO và không cần kết nối MQTT riêng, mà giao tiếp với backend qua HTTP và WebSocket.')
F('architecture.png','Hình 3.1 Sơ đồ kiến trúc và giao tiếp giữa các khối',15.8,16)
P('Trong chiều giám sát, cảm biến tạo dữ liệu, ESP32 publish các topic trạng thái, broker chuyển đến backend, backend ghi MySQL rồi phát JSON cho trình duyệt. Trong chiều điều khiển, dashboard gửi HTTP, backend kiểm tra yêu cầu rồi publish topic lệnh, ESP32 tiếp nhận và cập nhật trạng thái đầu ra. Topic /state được dùng để phản ánh kết quả ở mức phần mềm.',
'Mô hình đặt các dịch vụ xử lý trong LAN giúp quá trình thử nghiệm có thể quan sát trực tiếp. Broker được chạy bằng container Mosquitto; MySQL cung cấp dữ liệu lưu bền; Ollama chạy trên máy tính. Không có yêu cầu bắt buộc một dịch vụ cloud trong đường xử lý chính, nhưng toàn bộ chuỗi giám sát vẫn cần các thành phần cục bộ hoạt động đồng thời.')
H('Thiết kế phần cứng và phân bố chân')
T('Bảng 3.2 Phân bố chân theo cấu hình ESPHome',['Thiết bị','GPIO','Kiểu tín hiệu','Cấu hình quan trọng'],[
('DHT11','4','Dữ liệu vào và giao tiếp DHT','model DHT11; hai kênh nhiệt độ, độ ẩm'),('PIR HC-SR501','27','Ngõ vào số','on_press và on_release kích hoạt tự động'),('Module LDR','34','Ngõ vào số','Không khai báo inverted; chân chỉ nhập'),('LED chuyển động','25','Ngõ ra số','switch GPIO; có state_topic và command_topic'),('Relay phòng khách','26','Ngõ ra số','inverted: true; một kênh được cấu hình'),('Còi chủ động','33','Ngõ ra số','switch GPIO; còi tự động bật 500 ms')],[4.1,1.5,4.3,6.1])
P('DHT11 sử dụng một chân dữ liệu tại GPIO4 để cung cấp cả nhiệt độ và độ ẩm. Hai cảm biến nhị phân được đấu tới GPIO27 và GPIO34. Các khối có chung tham chiếu GND theo sơ đồ lắp của mô hình. Tín hiệu đưa về ESP32 phải phù hợp với mức logic của vi điều khiển; không suy ra mức điện áp ngõ ra của một module chỉ từ tên gọi nguồn cấp.',
'LED cần điện trở hạn dòng phù hợp. Với còi và module relay, việc chọn nguồn và tầng kích cần dựa trên dòng tiêu thụ thực tế của linh kiện. GPIO chỉ có vai trò phát tín hiệu điều khiển; khi triển khai với cơ cấu có dòng tải lớn hơn khả năng chân, cần tầng đệm phù hợp. Đối với tải điện lưới, phần tiếp điểm và dây dẫn phải được bố trí tách khỏi mạch tín hiệu thấp áp.',
'Ảnh mô hình cho thấy nhóm đã triển khai trên breadboard và dùng các dây nối để thuận tiện thay đổi cấu hình. Phiên bản đầu gồm ESP32, DHT11, PIR, LDR và relay. Phiên bản sau bổ sung LED, còi cùng chế độ thủ công. Việc chia phiên bản giúp thử riêng luồng đọc cảm biến trước khi mở rộng các hành động đầu ra.')
F(ROOT/'nha-thong-minh-dashboard/src/assets/deviceImage/version1.jpg','Hình 3.2 Mô hình lắp ráp phiên bản đầu',14.8,12)
F(ROOT/'nha-thong-minh-dashboard/src/assets/deviceImage/version2.jpg','Hình 3.3 Mô hình bổ sung LED và còi ở phiên bản sau',14.8,12)
H('Thiết kế firmware và chế độ hoạt động')
P('Thiết bị có tên cam-bien-phong-khach. Cấu hình khai báo logger, WiFi, MQTT, native API và OTA. MQTT discovery được tắt để sử dụng trực tiếp các topic do nhóm đặt. Một global kiểu bool có id che_do_thu_cong lưu cờ vận hành; restore_value: no và initial_value: false cho biết giá trị khởi tạo của biến này là chế độ tự động.',
'Công tắc template cong_tac_thu_cong nhận ON hoặc OFF qua thucong/set. Hai hành động turn_on_action và turn_off_action thay đổi global tương ứng. Tùy chọn optimistic giúp công tắc công bố trạng thái theo lệnh được xử lý, không dựa trên một cảm biến đo chế độ bên ngoài.',
'Ở nhánh on_press của PIR, chương trình chỉ tiếp tục khi cờ thủ công đang tắt. ESP32 bật LED, bật còi, chờ 500 ms rồi tắt còi. Sau bước này, chương trình kiểm tra binary_sensor ánh sáng; nếu trạng thái đang ON thì bật relay. Do vị trí lệnh trong YAML, relay được xét sau khoảng chờ của còi, chứ không được bật ngay ở đầu nhánh.',
'Ở nhánh on_release, nếu vẫn ở chế độ tự động, ESP32 tắt LED và relay. Cảm biến ánh sáng không có on_press hoặc on_release riêng để kích relay, nên việc ánh sáng thay đổi khi PIR vẫn giữ nguyên không tự tạo một lần đánh giá mới. Khi chuyển từ thủ công về tự động, cấu hình cũng chưa có hành động tính lại toàn bộ đầu ra ngay tại thời điểm chuyển chế độ.')
F('automation.png','Hình 3.4 Lưu đồ xử lý sự kiện PIR theo cấu hình hiện tại',15.4,19)
T('Bảng 3.3 Hành động tự động theo sự kiện',['Sự kiện','Điều kiện','Hành động'],[
('PIR chuyển ON','Thủ công OFF','Bật LED; còi bật 500 ms rồi tắt'),('Sau nhánh còi','Thủ công OFF và ánh sáng ON','Bật relay'),('PIR chuyển OFF','Thủ công OFF','Tắt LED và relay'),('PIR thay đổi','Thủ công ON','Bỏ qua nhánh tự động'),('Ánh sáng thay đổi','Mọi chế độ','Cập nhật trạng thái; không có nhánh relay riêng'),('Thủ công chuyển OFF','Lệnh người dùng','Đổi cờ; chờ các sự kiện tiếp theo')],[3.5,5,7.5])
P('Trạng thái ON của cảm biến ánh sáng trong firmware là mức logic dùng trong điều kiện tự động. Dashboard hiện đảo nhãn hiển thị ánh sáng để phù hợp cách nhóm quan sát module, nên cần phân biệt nhãn trên màn hình với giá trị MQTT gốc. Việc thống nhất cực tính ở firmware và giao diện sẽ giúp quy tắc điều khiển dễ giải thích hơn khi mở rộng mô hình.')
H('Thiết kế cây topic và payload MQTT')
P('Các topic được đặt theo cấu trúc phạm vi nhà, phòng, loại dữ liệu và hướng giao tiếp. Nhóm dùng chung tiền tố nha/phongkhach để backend có thể subscribe một nhánh. Topic cụ thể giúp tách riêng các kênh mà không cần đóng gói tất cả cảm biến vào một payload lớn.')
T('Bảng 3.4 Các topic trạng thái và lệnh',['Nhánh sau nha/phongkhach/','Payload hiện tại','Hướng và vai trò'],[
('nhietdo/state','Chuỗi số theo °C','ESP32 gửi nhiệt độ'),('doam/state','Chuỗi số theo %','ESP32 gửi độ ẩm tương đối'),('chuyendong/state','ON hoặc OFF','ESP32 gửi trạng thái PIR'),('anhsang/state','ON hoặc OFF','ESP32 gửi mức logic ánh sáng'),('relay/state và relay/set','ON hoặc OFF','Phản hồi relay và lệnh relay'),('den/state và den/set','ON hoặc OFF','Phản hồi LED và lệnh LED'),('coi/state và coi/set','ON hoặc OFF','Phản hồi còi và lệnh còi'),('thucong/state và thucong/set','ON hoặc OFF','Trạng thái và lệnh chọn chế độ')],[7.1,3.2,5.7])
P('Trong lịch sử còn tồn tại một số payload JSON tại den/state từ phiên bản khai báo LED trước đó. Với cấu hình switch hiện tại, payload dự kiến là ON hoặc OFF. Vì bảng dữ liệu lưu nguyên chuỗi nhận được, phần phân tích cần chuẩn hóa theo từng phiên bản thay vì coi mọi giá trị trong cùng topic có cùng định dạng.',
'Backend đang ghi mọi topic dưới nha/phongkhach/# vào cùng bảng cam_bien_data, bao gồm cả lệnh /set. Nhờ đó có thể đối chiếu một số lần gửi lệnh với phản hồi, nhưng tổng số hàng không bằng tổng số phép đo cảm biến. Việc tách một bảng lệnh riêng hoặc bổ sung loại bản tin sẽ giúp các truy vấn thống kê rõ nghĩa hơn.')
H('Thiết kế cơ sở dữ liệu')
P('Cơ sở dữ liệu nha_thong_minh có hai bảng: cam_bien_data và thoi_quen_nguoi_dung. Khóa chính id kiểu INT tự tăng nhận diện từng hàng. Hai bảng độc lập và không có quan hệ khóa ngoại trong cấu trúc hiện tại. Thời gian ở các trường DATETIME được tạo theo mặc định CURRENT_TIMESTAMP của MySQL khi thêm hàng.')
F('database.png','Hình 3.5 Cấu trúc hai bảng dữ liệu chính',15.6,10)
T('Bảng 3.5 Các trường trong bảng cam_bien_data',['Trường','Kiểu','Ràng buộc và ý nghĩa'],[
('id','INT','Khóa chính, AUTO_INCREMENT'),('topic','VARCHAR 255','NOT NULL; topic bản tin nhận được'),('gia_tri','VARCHAR 255','NOT NULL; payload dạng chuỗi'),('thoi_gian','DATETIME','Mặc định CURRENT_TIMESTAMP; thời điểm ghi nhận')],[3.2,4,8.8])
T('Bảng 3.6 Các trường trong bảng thoi_quen_nguoi_dung',['Trường','Kiểu','Ràng buộc và ý nghĩa'],[
('id','INT','Khóa chính, AUTO_INCREMENT'),('danh_muc','VARCHAR 50','NOT NULL; nhóm nội dung'),('noi_dung','VARCHAR 255','NOT NULL; mô tả thói quen'),('ngay_tao','DATETIME','Mặc định CURRENT_TIMESTAMP')],[3.2,4,8.8])
P('Thời gian MySQL trong lần kiểm tra sử dụng SYSTEM với múi giờ hệ điều hành SE Asia Standard Time. Khi xuất số liệu, nhóm đọc DATETIME ở dạng chuỗi để giữ nguyên giờ địa phương và tránh cộng múi giờ hai lần. Trên dashboard, các thời gian từ REST và thời gian tạo tại trình duyệt cần được thống nhất định dạng nếu muốn ghép dữ liệu theo cùng một trục thời gian.',
'Cấu trúc bảng hiện có khóa chính nhưng chưa có chỉ mục ghép theo topic và thoi_gian. Với số lượng bản ghi nhỏ, hệ thống vẫn truy vấn được lịch sử. Khi dữ liệu tăng, chỉ mục này có thể giúp các truy vấn WHERE topic và ORDER BY thoi_gian làm việc hiệu quả hơn; đồng thời cần quy định thời hạn lưu và cơ chế sao lưu dữ liệu.')
H('Xử lý dữ liệu tại backend')
P('Backend tạo một MQTT client kết nối mqtt://localhost:1883 và đăng ký nha/phongkhach/# sau khi kết nối thành công. Mỗi lần nhận dữ liệu, chương trình chuyển buffer thành chuỗi, thực hiện INSERT có tham số rồi gọi broadcast với JSON gồm topic và gia_tri. Dữ liệu được lưu trước khi gửi đến các WebSocket client đang ở trạng thái OPEN.',
'Cách dùng placeholder trong truy vấn INSERT giúp truyền payload qua tham số thay vì ghép trực tiếp vào chuỗi SQL. Hàm broadcast duyệt các kết nối hiện có và gửi cùng một JSON đến từng client còn mở. Dashboard ánh xạ topic sang khóa state tương ứng để cập nhật các thẻ và chuỗi điểm dữ liệu.',
'Do việc broadcast đặt sau câu lệnh await ghi dữ liệu, lỗi MySQL có thể làm gián đoạn nhánh cập nhật giao diện. Backend hiện cần bổ sung xử lý lỗi riêng cho luồng MQTT và thông báo trạng thái dịch vụ để phản ánh tình huống cơ sở dữ liệu tạm thời không sẵn sàng. Đây là cải tiến về độ ổn định khi chuyển từ mô hình thử nghiệm sang vận hành dài ngày.')
H('Thiết kế các REST API')
P('Backend sử dụng cổng 3000, cho phép parse JSON và nhận các yêu cầu từ dashboard. Các API truy vấn trả JSON để giao diện có thể dùng trực tiếp. Với lịch sử, đường dẫn :loai được ánh xạ qua danh sách định sẵn để tránh chọn topic tùy ý từ đầu vào người dùng.')
T('Bảng 3.7 Các API của hệ thống',['Phương thức và đường dẫn','Đầu vào chính','Kết quả'],[
('GET /api/lichsu/:loai','nhietdo, doam, chuyendong, anhsang, relay','Tối đa 50 hàng mới nhất, gồm gia_tri và thoi_gian'),('POST /api/relay','trang_thai: ON hoặc OFF','Publish relay/set; trả thông báo đã gửi lệnh'),('POST /api/che-do-thu-cong','trang_thai: ON hoặc OFF','Publish thucong/set'),('GET /api/bat-thuong','Không có','Danh sách mô tả bất thường và ghi chú nếu chưa đủ thời gian'),('GET /api/thoi-quen','Không có','Danh sách nhóm theo danh_muc'),('POST /api/thoi-quen','danh_muc, noi_dung','Thêm hàng và trả id, danh mục, nội dung'),('DELETE /api/thoi-quen/:id','Mã bản ghi','Xóa mục thói quen theo id'),('POST /api/chatbot','cau_hoi','tra_loi hoặc thông báo lỗi')],[6.3,4.1,5.6])
P('Khi nhận lệnh relay hoặc chế độ thủ công, backend kiểm tra giá trị chỉ được là ON hoặc OFF; đầu vào khác trả HTTP 400. Với loại lịch sử không có trong ánh xạ, backend trả HTTP 404. API thêm thói quen kiểm tra không được để trống và giới hạn độ dài tên danh mục là 50 ký tự.',
'Phản hồi HTTP thành công của API điều khiển chỉ xác nhận backend đã xử lý yêu cầu gửi lệnh. Chương trình không chờ một thông điệp xác nhận riêng có mã giao dịch từ ESP32. Vì thế, việc hiển thị kết quả cần dựa thêm vào topic trạng thái; để xác minh tác động vật lý còn cần phản hồi từ mạch đo hoặc quan sát tại mô hình.')
H('Gợi ý môi trường và phát hiện bất thường')
P('Hàm kiemTraGoiY áp dụng các điều kiện cố định cho nhiệt độ và độ ẩm. Mỗi khi nhận bản tin, backend tạo gợi ý và so sánh nội dung với lần trước trên cùng topic. Nếu nội dung thay đổi, chương trình phát JSON type: canh_bao qua WebSocket. Cơ chế này giảm việc gửi lặp cùng một câu, nhưng khi giá trị số thay đổi, nội dung câu cũng thay đổi dù vẫn thuộc một mức cảnh báo.')
T('Bảng 3.8 Các ngưỡng gợi ý trong mã nguồn',['Đại lượng','Điều kiện','Phản hồi'],[
('Nhiệt độ','T > 33°C','canh_bao; gợi ý quạt hoặc điều hòa'),('Nhiệt độ','30°C < T ≤ 33°C','goi_y; cân nhắc bật quạt'),('Nhiệt độ','T ≤ 30°C','Không tạo gợi ý nhiệt độ'),('Độ ẩm','H < 40%','goi_y; độ ẩm thấp'),('Độ ẩm','H > 80%','goi_y; độ ẩm cao'),('Độ ẩm','40% ≤ H ≤ 80%','Không tạo gợi ý độ ẩm')],[3.1,5.1,7.8])
P('API bất thường dùng thời gian từ bản ghi đầu tiên đến hiện tại để quyết định có đủ tối thiểu 3 ngày hay chưa. Sau đó, chương trình tính nhiệt độ trung bình của các hàng trong 7 ngày gần nhất và so với nhiệt độ mới nhất; nếu độ lệch tuyệt đối lớn hơn 3°C thì thêm thông báo. Với chuyển động, chương trình tìm xem trong khung giờ hiện tại có bản tin ON hôm nay nhưng không có bản tin ON nào ở cùng giờ trong các ngày trước trong cửa sổ 7 ngày hay không.',
'Đây là các quy tắc thống kê đơn giản, không phải mô hình học máy được huấn luyện. Điều kiện đủ 3 ngày hiện tại đo khoảng cách từ bản ghi đầu đến thời điểm chạy, không kiểm tra có đủ 3 ngày dữ liệu thực sự. Khi dữ liệu không liên tục hoặc bị phát lại, kết quả cần được hiểu theo đặc điểm tập dữ liệu đang có. Nhóm chưa dùng đầu ra này để tự động can thiệp vào thiết bị.')
H('Thiết kế trợ lý AI')
P('Khi nhận cau_hoi, backend gọi phatHienLenh trước. Hàm chuyển chuỗi về chữ thường rồi kiểm tra các mẫu tiếng Việt và tiếng Anh cho bốn thao tác: bật đèn, tắt đèn, bật còi và tắt còi. Nếu khớp, backend publish thucong/set bằng ON, sau đó publish lệnh tới den/set hoặc coi/set và trả lời ngay bằng một câu xác nhận đã thực hiện nhánh điều khiển.',
'Trong ngữ cảnh này, từ “đèn” trong nhánh chatbot được ánh xạ tới LED GPIO25, không phải relay GPIO26. Relay có API và nút điều khiển riêng. Cách ánh xạ này cần được trình bày rõ cho người dùng để tránh nhầm giữa đèn chỉ báo trên mô hình với một thiết bị chiếu sáng nối qua relay.',
'Với câu hỏi không khớp lệnh, hàm taoTomTat lấy nhiệt độ, độ ẩm mới nhất; tạo gợi ý; kiểm tra bất thường; đọc thói quen. Backend ghép những nội dung này vào system message và gọi Ollama. Mỗi yêu cầu chỉ gửi system message và câu hỏi hiện tại; các tin nhắn cũ đang hiển thị trên giao diện không được chuyển thành lịch sử hội thoại đầy đủ cho mô hình.')
F('chatbot-flow.png','Hình 3.6 Luồng nhận diện lệnh và hỏi đáp của chatbot',15.6,17)
P('Mô hình được yêu cầu trả lời tiếng Anh trong 2–4 câu ngắn. Nhóm có thể đặt câu hỏi về dữ liệu cảm biến, các gợi ý, bất thường hoặc nội dung thói quen đã ghi. Chất lượng câu trả lời phụ thuộc vào ngữ cảnh được tạo và trạng thái dịch vụ Ollama; dữ liệu cảm biến được giữ trong ngữ cảnh nhưng mô hình không tự kiểm chứng tuổi của bản tin.',
'Regex giúp triển khai nhanh các câu lệnh đã xác định. Khi câu chứa phủ định, nhiều thiết bị hoặc nhiều hành động, một mẫu đơn giản có thể nhận diện không như ý định. Hướng hoàn thiện là chuẩn hóa tập lệnh, kiểm tra đầu vào và hiển thị rõ trạng thái đã gửi lệnh, đã nhận phản hồi thiết bị, thay vì dùng một câu khẳng định chung cho tất cả các bước.')
H('Thiết kế giao diện giám sát')
P('Dashboard được chia thành bốn trang chức năng: bảng điều khiển, thiết bị sử dụng, thói quen và trợ lý AI. Bảng điều khiển có bốn thẻ cảm biến, khối relay, trạng thái chế độ thủ công, biểu đồ nhiệt độ và độ ẩm, ô hỏi đáp nhanh cùng các ghi chú. Trang thiết bị liệt kê linh kiện và ảnh mô hình theo phiên bản để giải thích cấu trúc sản phẩm.',
'Khi trang được mở, dashboard gọi API lịch sử cho năm loại nhietdo, doam, chuyendong, anhsang và relay. Giá trị đầu tiên trong kết quả được dùng làm trạng thái gần nhất. Hai chuỗi nhiệt độ, độ ẩm được đảo về thứ tự thời gian để ghép vào biểu đồ, sau đó giữ tối đa 50 mốc dữ liệu. Người dùng có thể chọn 20 hoặc 50 điểm hiển thị.',
'Kết nối WebSocket cập nhật state khi nhận bản tin. Giá trị relay được dùng để bỏ trạng thái pending của nút điều khiển; các cảnh báo được hiển thị trong danh sách tối đa 5 mục gần nhất. Nếu kết nối đóng, dashboard thử kết nối lại sau 3 giây. Khi component được hủy, bộ hẹn giờ và socket được đóng để tránh tiếp tục nhận dữ liệu ngoài vòng đời giao diện.')
F('dashboard-status.png','Hình 3.7 Giao diện trạng thái khi sử dụng dữ liệu đã lưu',15.8,8)
P('Ảnh giao diện thể hiện dữ liệu tham chiếu từ các bản tin đã lưu: nhiệt độ 29,9°C, độ ẩm 83%, PIR OFF và relay OFF. Nhãn ánh sáng được hiển thị qua ánh xạ đảo của frontend. Thời gian ở các thẻ trong ảnh là thời điểm dữ liệu được backend nhận lại; đây là minh họa hoạt động hiển thị bằng dữ liệu lịch sử.',
'Trạng thái chế độ thủ công chưa có API lịch sử tương ứng trong ánh xạ :loai. Khi tải trang, frontend có thể chưa biết trạng thái này cho đến khi nhận thucong/state qua WebSocket, trong khi giao diện hiện diễn giải giá trị chưa biết như OFF. Cần có trạng thái “chưa xác định” và cơ chế tải ban đầu để đồng bộ chính xác sau khi trình duyệt mở lại.')

chapter('CHƯƠNG 4 KẾT QUẢ THỬ NGHIỆM VÀ ĐÁNH GIÁ')
H('Môi trường và phạm vi đánh giá')
P('Các thử nghiệm trước của nhóm được thực hiện trên mô hình ESP32 với các cảm biến và cơ cấu chấp hành đã trình bày. Nhóm đã chạy thành công các chức năng đọc cảm biến, truyền MQTT, giám sát trên web, điều khiển tự động và thủ công, lưu lịch sử, quản lý thói quen và hỏi đáp AI. Dữ liệu còn lưu trong MySQL được sử dụng để thống kê và minh họa kết quả.',
'Ngày 15/10/2026, các dịch vụ phần mềm được khởi động lại để kiểm tra khả năng truy cập dữ liệu và phản hồi. Broker Mosquitto, backend, dashboard và Ollama được chạy; MySQL sử dụng phiên bản 8.0.44. Phần đánh giá ở lần khởi động lại gồm API đọc lịch sử, đầu vào không hợp lệ, kết nối WebSocket, các hàm quy tắc và chatbot. Việc gửi thêm các lệnh có thể làm thay đổi cơ cấu chấp hành không được dùng làm nguồn thống kê cho dữ liệu lịch sử.',
'Nhóm sử dụng lại trạng thái cũ để trình bày giao diện khi cần minh họa. Hai phiên dữ liệu chính ngày 01/09/2026 và 08/09/2026 được phân tích độc lập với bản tin retained nhận lại vào tháng 10. Số lượng các hàng mô tả dữ liệu lưu trữ, không được chuyển thành tỷ lệ thành công hoặc độ chính xác cảm biến nếu chưa có quy trình đo tương ứng.')
T('Bảng 4.1 Các nguồn dùng để đánh giá',['Nguồn','Nội dung','Cách diễn giải'],[
('Thử nghiệm mô hình của nhóm','Các chức năng cảm biến và điều khiển đã chạy thành công','Kết quả chức năng của các lần chạy trước'),('MySQL lịch sử','2.077 hàng trong hai ngày 01/09 và 08/09/2026','Thống kê bản tin, phạm vi giá trị và ví dụ lệnh phản hồi'),('Cấu hình và mã nguồn','YAML ESPHome, backend và frontend','Xác định thuật toán, điều kiện, topic và API'),('Kiểm tra ngày 15/10/2026','MQTT, REST, WS, logic, build, chatbot','Khả năng hoạt động phần mềm trong lần kiểm tra lại')],[4,6.3,5.7])
H('Kết quả thống kê dữ liệu lịch sử')
P('Tập dữ liệu lịch sử bắt đầu lúc 13:35:51 ngày 01/09/2026 và kết thúc lúc 16:10:37 ngày 08/09/2026 theo giờ Việt Nam. Có 849 hàng trong ngày 01/09 và 1.228 hàng trong ngày 08/09, tổng cộng 2.077 hàng. Dữ liệu nằm trong hai phiên chạy; khoảng cách bảy ngày giữa mốc đầu và mốc cuối không thể được hiểu là bảy ngày giám sát liên tục.',
'Trong tổng số hàng, 1.950 hàng có topic kết thúc bằng /state, còn 127 hàng có topic kết thúc bằng /set. Bản tin chuyển động có số lượng lớn nhất vì được gửi theo sự kiện, trong khi DHT11 cung cấp hai kênh theo chu kỳ. Các bản tin khởi tạo, retained hoặc công bố lại trạng thái cũng có thể góp phần vào số lượng hàng.')
hist={x['topic'].split('/')[2]+'/'+x['topic'].split('/')[3]:x['n'] for x in E['queries']['historical_topics']}
T('Bảng 4.2 Số bản ghi theo topic trong hai phiên lịch sử',['Topic sau nha/phongkhach/','Số hàng','Loại'],[(k,str(v),'Trạng thái' if k.endswith('/state') else 'Lệnh') for k,v in hist.items()]+[('Tổng cộng','2.077','1.950 trạng thái và 127 lệnh')],[8,2.6,5.4])
P('Có 10 mục thói quen thuộc 5 danh mục trong cơ sở dữ liệu. Các danh mục gồm Buổi tối, Thường hay quên, Sử dụng quạt, Sử dụng thiết bị và Ánh sáng. Chức năng này cung cấp nội dung tham chiếu cho chatbot; việc có một câu mô tả giờ bật thiết bị trong bảng không đồng nghĩa một lịch tự động đã được tạo.')
H('Đặc điểm số liệu nhiệt độ và độ ẩm')
T('Bảng 4.3 Thống kê mô tả nhiệt độ và độ ẩm',['Đại lượng','Số hàng','Nhỏ nhất','Lớn nhất','Trung bình'],[
('Nhiệt độ','304','29,4°C','35,0°C','31,251°C'),('Độ ẩm tương đối','303','54%','88%','69,033%')],[4.4,2.3,3.1,3.1,3.1])
P('Trung bình trong bảng được tính trên toàn bộ các hàng tương ứng trong hai phiên lịch sử, với mỗi hàng có cùng trọng số. Đây là trung bình các mẫu đã ghi, không phải trung bình theo thời lượng hoạt động của cả khoảng ngày. Số hàng nhiệt độ và độ ẩm khác nhau một đơn vị nên cũng không thể giả định hai đại lượng luôn có đầy đủ cặp tại mọi thời điểm.',
'Khoảng nhiệt độ đi qua cả ngưỡng gợi ý trên 30°C và ngưỡng cảnh báo trên 33°C. Độ ẩm lớn nhất vượt ngưỡng 80% được khai báo trong backend. Vì vậy, dữ liệu đã lưu cung cấp ví dụ phù hợp để minh họa các mức gợi ý của hệ thống. Không có thiết bị đo chuẩn đi kèm trong tập dữ liệu để tính sai số tuyệt đối hay sai số phần trăm của DHT11.',
'Giá trị nhiệt độ và độ ẩm ở cuối phiên ngày 08/09 là 29,9°C và 83%. Khi broker được bật lại vào ngày 15/10, các giá trị này còn trong các bản tin retained và được backend nhận lại. Chúng được dùng để hiển thị giao diện và kiểm tra hỏi đáp theo dữ liệu đã có, thay vì coi đó là một cặp đo mới tại phòng khách ngày 15/10.')
H('Minh họa luồng lệnh và phản hồi từ dữ liệu cũ')
P('Trong lịch sử ngày 08/09 có các chuỗi lệnh LED và relay kèm trạng thái trả về. Ví dụ, lúc 16:02:13 backend ghi nhận thucong/set bằng ON và den/set bằng ON, sau đó den/state bằng ON. Lúc 16:02:24 xuất hiện den/set bằng OFF và den/state bằng OFF. Các hàng này phù hợp với luồng điều khiển thủ công qua MQTT đã thiết kế.')
T('Bảng 4.4 Một số sự kiện lệnh và trạng thái ngày 08 tháng 09',['Giờ địa phương','Topic sau nha/phongkhach/','Payload','Ý nghĩa'],[
('16:02:13','thucong/set','ON','Yêu cầu chế độ thủ công'),('16:02:13','den/set','ON','Lệnh bật LED'),('16:02:13','den/state','ON','Trạng thái LED được công bố'),('16:02:24','den/set','OFF','Lệnh tắt LED'),('16:02:24','den/state','OFF','Trạng thái LED được công bố'),('16:03:18','relay/set','ON','Lệnh bật relay'),('16:03:18','relay/state','ON','Trạng thái relay được công bố'),('16:03:27','relay/set','OFF','Lệnh tắt relay'),('16:03:27','relay/state','OFF','Trạng thái relay được công bố')],[3,5.9,1.8,5.3])
P('Thời gian trong bảng MySQL có độ phân giải đến giây và là thời điểm ghi tại máy chủ. Nhiều lệnh và phản hồi nằm trong cùng một giây nên bảng chỉ chứng minh trật tự bản ghi, không cho phép kết luận độ trễ bằng 0 ms. Ngoài ra, trạng thái GPIO hoặc relay ở mức phần mềm không thay thế phép đo điện trên tải ngoài.',
'Một số lệnh relay được ghi lặp hai lần trong cùng giây. Điều này cho thấy việc đếm số hàng /set chưa đủ để suy ra số thao tác duy nhất của người dùng. Để đo tỷ lệ xác nhận và độ trễ điều khiển một cách chắc chắn hơn, phiên bản sau cần mã lệnh, thời điểm gửi, thời điểm thiết bị nhận và phản hồi tương ứng.')
H('Kết quả thử nghiệm chức năng của mô hình')
P('Nhóm đã thực hiện thành công các kịch bản chức năng chính trong các lần chạy trước. Bảng dưới tổng hợp theo từng chức năng; phần số liệu và các bản tin minh họa ở trên bổ sung bằng chứng cho lưu trữ, trạng thái và một số thao tác điều khiển. Bảng không biểu diễn một thử nghiệm thống kê với số lần lặp cố định.')
T('Bảng 4.5 Tổng hợp các kịch bản chức năng đã thử nghiệm',['Kịch bản','Kết quả thực hiện','Nhận xét'],[
('Đọc DHT11','Thành công','Có dữ liệu nhiệt độ và độ ẩm trong MySQL'),('Đọc PIR và LDR','Thành công','Có bản tin ON/OFF theo từng topic'),('PIR kích LED và còi','Thành công','Đã chạy mô hình; cấu hình còi có khoảng chờ 500 ms'),('Relay tự động','Thành công','Thực hiện theo trạng thái PIR, ánh sáng và chế độ'),('Bật tắt relay từ dashboard','Thành công','Có lệnh và trạng thái relay trong lịch sử'),('Điều khiển LED, còi bằng câu lệnh','Thành công','Có den/set, coi/set và các trạng thái tương ứng'),('Chuyển chế độ thủ công','Thành công','Có thucong/set và thucong/state'),('Lưu và đọc lịch sử','Thành công','Các phiên dữ liệu được giữ trong MySQL'),('Quản lý thói quen','Thành công','Có 10 mục thuộc 5 danh mục'),('Hỏi đáp AI','Thành công','Trả lời theo dữ liệu cảm biến và ngữ cảnh')],[5.3,3.1,7.6])
H('Kiểm tra lại phần mềm ngày 15 tháng 10 năm 2026')
P('Năm API lịch sử đều trả HTTP 200 và 50 hàng mới nhất cho từng loại hỗ trợ. Loại không có trong ánh xạ trả HTTP 404. Hai API điều khiển relay và chế độ thủ công trả HTTP 400 khi nhận trạng thái INVALID, phù hợp với ràng buộc ON/OFF. Kiểm tra đầu vào không hợp lệ không tạo thêm lệnh điều khiển.',
'Kết nối WebSocket tới cổng 8081 được mở thành công. Một MQTT client quan sát nhận 8 topic trạng thái với cờ retain bằng true và QoS bằng 0. Backend đồng thời nhận lại các trạng thái được broker giữ, làm tăng cơ sở dữ liệu từ 2.077 lên 2.085 hàng trong lần khởi động. Tám hàng này được loại khỏi thống kê hai phiên lịch sử.',
'Quá trình npm run build hoàn thành thành công, tạo gói giao diện trong dist. Công cụ build thông báo một gói JavaScript sau tối ưu lớn hơn ngưỡng khuyến nghị 500 kB; đây là hướng để tối ưu tải trang bằng chia nhỏ mã khi mở rộng ứng dụng, không làm thất bại quá trình build hiện tại.')
T('Bảng 4.6 Kết quả kiểm tra dịch vụ phần mềm',['Hạng mục','Kết quả quan sát','Đánh giá'],[
('MQTT cổng 1883','Kết nối thành công; nhận 8 trạng thái retained','Đạt kết nối và nhận trạng thái đã lưu'),('5 API lịch sử','HTTP 200; mỗi loại trả 50 hàng','Đạt'),('Loại lịch sử không tồn tại','HTTP 404','Đạt kiểm tra đầu vào'),('Trạng thái điều khiển không hợp lệ','HTTP 400 ở hai API','Đạt kiểm tra đầu vào'),('API thói quen','HTTP 200; 10 mục, 5 danh mục','Đạt đọc dữ liệu'),('API bất thường','HTTP 200; danh sách rỗng ở lần gọi','Đạt phản hồi; cần đọc theo cửa sổ dữ liệu'),('WebSocket cổng 8081','Mở kết nối thành công','Đạt kết nối trình duyệt và backend'),('Build dashboard','Hoàn thành; có khuyến nghị kích thước gói','Đạt'),('Chatbot hỏi dữ liệu','HTTP 200; trả đúng 29,9°C và 83%','Đạt trong lần kiểm tra lại')],[4.8,6.8,4.4])
H('Kiểm tra các ngưỡng và nhận diện câu lệnh')
P('Nhóm đối chiếu trực tiếp hai hàm kiemTraGoiY và phatHienLenh trong backend. Các phép kiểm tra chạy tách biệt phần kết nối phần cứng, do đó không phát lệnh MQTT và không làm thay đổi đầu ra. Tổng cộng 15 trường hợp có kết quả đúng theo điều kiện đã viết trong mã nguồn.')
T('Bảng 4.7 Các trường hợp kiểm tra logic',['Đầu vào','Kết quả mong đợi và quan sát'],[
('Nhiệt độ 30°C','Không tạo gợi ý'),('Nhiệt độ 30,1°C','goi_y'),('Nhiệt độ 33°C','goi_y'),('Nhiệt độ 33,1°C','canh_bao'),('Độ ẩm 39,9%','goi_y'),('Độ ẩm 40%','Không tạo gợi ý'),('Độ ẩm 80%','Không tạo gợi ý'),('Độ ẩm 80,1%','goi_y'),('bật đèn','den: ON'),('tắt đèn','den: OFF'),('bật còi','coi: ON'),('tắt còi','coi: OFF'),('turn on the light','den: ON'),('turn off the buzzer','coi: OFF'),('Nhiệt độ hiện tại là bao nhiêu?','Không nhận là lệnh điều khiển')],[8,8])
P('Các giá trị đúng bằng 30°C, 33°C, 40% và 80% giúp xác định rõ dấu lớn hơn hoặc nhỏ hơn trong điều kiện. Kết quả 15/15 chỉ phản ánh các trường hợp cụ thể đã kiểm tra. Nó không chứng minh mọi câu tiếng Việt và tiếng Anh đều được hiểu đúng, cũng không thay cho phép kiểm tra toàn bộ đường điều khiển đến tải vật lý.')
H('Kết quả hỏi đáp AI')
P('Trong các lần chạy trước, chatbot đã trả lời tốt theo xác nhận của nhóm. Khi kiểm tra lại, câu hỏi “What are the temperature and humidity readings? Answer in one sentence.” nhận phản hồi HTTP 200 với nội dung “The current temperature is 29.9°C, and the humidity is at 83%.” Hai giá trị trùng với các giá trị trạng thái đã được nạp vào backend.',
'Một lần gọi thành công trong phiên kiểm tra có thời gian chờ khoảng 19,2 giây ở phía chương trình gọi API. Đây là thời gian của một lần hỏi đáp cụ thể, không phải độ trễ trung bình, độ trễ điều khiển hoặc cam kết hiệu năng của hệ thống. Dịch vụ cần thời gian nạp mô hình khi khởi động; một yêu cầu ban đầu chưa thành công và lần kiểm tra lại sau khi mô hình sẵn sàng đã trả lời đúng.',
'Kết quả cho thấy nhánh lấy ngữ cảnh từ MySQL và gọi Ollama hoạt động. Từ “current” trong câu trả lời hiện do hướng dẫn ngữ cảnh của backend tạo ra; trong trường hợp này, dữ liệu là trạng thái cũ được dùng để minh họa. Hướng hoàn thiện là đưa thời gian đo và độ mới của dữ liệu vào ngữ cảnh để trợ lý trả lời “giá trị ghi nhận gần nhất” khi không có phép đo mới.')
H('Ưu điểm của hệ thống')
P('Hệ thống tạo được một chuỗi IoT hoàn chỉnh từ cảm biến vật lý tới giao diện web và dữ liệu lưu bền. Mỗi khối có vai trò rõ ràng, giúp nhóm quan sát và kiểm tra sự cố theo đường truyền. Cây topic MQTT và các API nhỏ cho phép bổ sung thành phần mà không phải thay đổi đồng thời toàn bộ kiến trúc.',
'Logic phản ứng với PIR được đặt tại ESP32 nên có thể thực hiện ở thiết bị khi không chờ backend quyết định. Dữ liệu lịch sử được giữ trong MySQL và có thể sử dụng lại khi phần mềm khởi động. Dashboard kết hợp truy vấn ban đầu và cập nhật WebSocket, hỗ trợ cả theo dõi và tương tác trong cùng một giao diện.',
'Trợ lý AI bổ sung hình thức khai thác dữ liệu bằng ngôn ngữ tự nhiên. Các lệnh đơn giản đi qua nhánh xác định trước, trong khi phần hỏi đáp dùng một mô hình cục bộ có ngữ cảnh. Cách kết hợp này phù hợp với phạm vi học phần và giúp nhóm nhận diện rõ hơn sự khác nhau giữa điều khiển theo quy tắc, thống kê đơn giản và hỏi đáp bằng LLM.')
H('Giới hạn và hướng hoàn thiện')
P('Về đồng bộ trạng thái, payload hiện chưa có timestamp từ ESP32, mã thiết bị, số thứ tự hoặc cờ dữ liệu cũ trong bản ghi MySQL. Backend cũng chưa lưu retain của MQTT. Vì vậy, dữ liệu phát lại có thể được ghi với thời gian nhận mới. Cần bổ sung các trường nguồn gốc này và hiển thị tuổi dữ liệu để người dùng phân biệt trạng thái tham khảo với trạng thái vừa đo.',
'Về điều khiển, chế độ tự động được kích theo cạnh PIR, chưa tính lại đầu ra khi chỉ ánh sáng đổi hoặc khi vừa thoát thủ công. Relay được xét sau khoảng chờ 500 ms của còi. Giao diện cần đồng bộ giá trị thủ công khi khởi tạo và thống nhất ánh xạ cực tính của LDR. Các thay đổi này sẽ làm hành vi hệ thống dễ dự đoán hơn trong nhiều tình huống.',
'Về dữ liệu và thống kê, điều kiện đủ 3 ngày của phần bất thường chưa kiểm tra số ngày có mẫu thực sự; cửa sổ trung bình có thể chỉ chứa dữ liệu phát lại. Chưa có chỉ số đo độ đầy đủ dữ liệu hoặc cơ chế hiệu chuẩn. Khi đánh giá phiên bản tiếp theo, cần có các phiên đo liên tục, nhật ký điều kiện thử và thiết bị tham chiếu nếu muốn kết luận về độ chính xác.',
'Về độ ổn định và truy cập, backend cần xử lý lỗi cho từng luồng, kiểm tra dữ liệu đầu vào đầy đủ hơn và giới hạn phạm vi CORS. Các giao thức đang dùng mqtt://, http:// và ws:// trong LAN; xác thực API và mã hóa đường truyền cần được bổ sung trước khi cho phép truy cập từ môi trường rộng hơn. Broker và máy chủ hiện cũng là các điểm phụ thuộc chung của hệ thống.',
'Về trợ lý, cần bổ sung timestamp, lịch sử hội thoại có giới hạn và trạng thái dịch vụ mô hình. Regex nên xử lý phủ định, nhiều hành động và các câu không chắc chắn. Nội dung thói quen chỉ nên trở thành lịch tự động khi có bộ lập lịch và quy tắc xác nhận riêng; mô hình hiện tại chưa thực hiện bước này.')
T('Bảng 4.8 Các hướng phát triển ưu tiên',['Ưu tiên','Nội dung','Kết quả mong muốn'],[
('1','Chuẩn hóa payload, thời gian đo và retain','Hiển thị đúng độ mới, phân biệt dữ liệu phát lại'),('2','Đồng bộ chế độ và cực tính ánh sáng','Giảm khác biệt giữa firmware, MQTT và dashboard'),('3','Bổ sung xác nhận lệnh và timeout','Phân biệt lệnh đã gửi, đã nhận, đã thực thi'),('4','Xử lý lỗi và chỉ mục dữ liệu','Vận hành ổn định hơn khi dữ liệu tăng'),('5','Xác thực API và mã hóa phù hợp','Kiểm soát quyền truy cập khi mở rộng'),('6','Ngữ cảnh AI có thời gian, hội thoại','Trả lời rõ nguồn dữ liệu và bám ngữ cảnh'),('7','Nhiều phòng và bộ lập lịch','Mở rộng ứng dụng sau khi các khối cơ bản ổn định')],[1.9,7.1,7])

chapter('KẾT LUẬN')
P('Đề tài đã xây dựng mô hình giám sát và điều khiển nhà thông minh cho phòng khách, kết hợp ESP32, DHT11, PIR, module ánh sáng, relay, LED và còi với các dịch vụ MQTT, Node.js, MySQL và dashboard React. Hệ thống thể hiện được hai luồng chính của một ứng dụng IoT: thu thập trạng thái từ thiết bị lên ứng dụng và gửi lệnh từ ứng dụng về cơ cấu chấp hành.',
'Các chức năng tự động và thủ công đã được nhóm thử nghiệm thành công. Dữ liệu lịch sử gồm 2.077 hàng cung cấp cơ sở để mô tả phạm vi hoạt động, giá trị môi trường và ví dụ về lệnh, trạng thái phản hồi. Kiểm tra lại các dịch vụ xác nhận khả năng truy vấn lịch sử, kết nối MQTT, WebSocket, build giao diện và hỏi đáp AI; các phép kiểm tra logic tại ngưỡng cho kết quả đúng với mã nguồn.',
'Phần trợ lý AI sử dụng dữ liệu của hệ thống làm ngữ cảnh để trả lời câu hỏi, đồng thời tách một số lệnh đèn, còi sang nhánh quy tắc. Đây là bước mở rộng giao diện tương tác, chưa phải mô hình tự học hành vi và tự quyết định toàn bộ hoạt động trong nhà. Việc lưu thói quen hiện phục vụ tham chiếu cho hỏi đáp.',
'Qua quá trình thực hiện, nhóm vận dụng được kiến thức phần cứng, lập trình cấu hình thiết bị, mạng máy tính, cơ sở dữ liệu và phát triển web. Các nội dung cần tiếp tục hoàn thiện là đồng bộ thời gian và độ mới dữ liệu, thống nhất trạng thái giữa thiết bị với dashboard, bổ sung xác nhận lệnh, xử lý lỗi, kiểm soát truy cập và kiểm thử dài ngày. Những cải tiến này là cơ sở để mở rộng từ mô hình học tập sang một hệ thống có phạm vi sử dụng lớn hơn.')

chapter('TÀI LIỆU THAM KHẢO')
refs=[
('[1] Espressif Systems','GPIO and RTC GPIO, ESP-IDF Programming Guide','https://docs.espressif.com/projects/esp-idf/en/stable/esp32/api-reference/peripherals/gpio.html'),
('[2] ESPHome','MQTT Client Component','https://esphome.io/components/mqtt/'),
('[3] ESPHome','DHT Temperature and Humidity Sensor','https://esphome.io/components/sensor/dht/'),
('[4] MQTT.org','MQTT The Standard for IoT Messaging','https://mqtt.org/'),
('[5] Express','Routing','https://expressjs.com/en/guide/routing/'),
('[6] React','Quick Start','https://react.dev/learn'),
('[7] Ollama','Generate a chat message','https://docs.ollama.com/api/chat'),
]
for author,title,url in refs:
    p=D.add_paragraph(f'{author}. {title}. Tài liệu trực tuyến. Truy cập ngày 15/10/2026. {url}');p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(10)
P('[8] Nhóm sinh viên D18CNPM5. Mã nguồn dự án nhà thông minh: nha-thong-minh-backend/index.js, nha-thong-minh-dashboard/src/App.jsx; cấu hình ESPHome và ảnh mô hình. Tài liệu nội bộ của đề tài, 2026.',
'[9] Cơ sở dữ liệu nha_thong_minh. Hai phiên dữ liệu ngày 01/09/2026 và 08/09/2026; kết quả kiểm tra dịch vụ ngày 15/10/2026. Dữ liệu thực nghiệm nội bộ của nhóm.')

chapter('PHỤ LỤC A CẤU HÌNH ESPHOME')
P('Cấu hình dưới đây thể hiện các chân, topic và hành động của phiên bản mô hình sử dụng trong báo cáo. Các giá trị xác thực được tham chiếu qua secrets.yaml. Việc tách secrets không thay đổi thuật toán PIR, LED, relay, còi và chế độ thủ công.')
code((WORK/'firmware-report.yaml').read_text(encoding='utf8'))

chapter('PHỤ LỤC B HƯỚNG DẪN CHẠY VÀ ĐỐI CHIẾU')
H('Khởi động các dịch vụ')
P('Kiểm tra MySQL80 đang chạy và cơ sở dữ liệu nha_thong_minh có hai bảng cam_bien_data, thoi_quen_nguoi_dung. Kiểm tra file .env của backend chứa cấu hình kết nối DB_HOST, DB_USER, DB_PASSWORD và DB_NAME. Không đưa giá trị mật khẩu vào log hoặc tài liệu chia sẻ.')
code('docker start mosquitto\n\n# Tai thu muc nha-thong-minh-backend\nnode index.js\n\n# Tai thu muc nha-thong-minh-dashboard\nnpm run dev\n\n# Dich vu AI\nollama serve\nollama list')
P('Trên Windows PowerShell, nếu npm.ps1 bị chính sách thực thi chặn, dùng npm.cmd cho cùng lệnh. Khi Ollama đã chạy thì không cần mở thêm một tiến trình serve. Danh sách model cần có llama3.1:latest để phù hợp với cấu hình gọi hiện tại.')
H('Địa chỉ dùng khi chạy trên máy chủ')
T(None,['Dịch vụ','Địa chỉ hoặc cổng'],[('Dashboard phát triển','http://localhost:5173'),('REST API','http://localhost:3000'),('WebSocket','ws://localhost:8081'),('MQTT backend','mqtt://localhost:1883'),('Ollama','http://localhost:11434'),('MQTT phía ESP32','192.168.0.102:1883 trong cấu hình của nhóm')],[5,11])
H('Ví dụ nội dung yêu cầu và phản hồi')
P('Các ví dụ JSON dưới mô tả giao ước API. Yêu cầu relay có thể làm thay đổi thiết bị khi hệ thống đang hoạt động; chỉ gửi khi chủ động kiểm tra luồng điều khiển.')
code('POST /api/relay\n{"trang_thai":"ON"}\n\nPOST /api/che-do-thu-cong\n{"trang_thai":"OFF"}\n\nPOST /api/chatbot\n{"cau_hoi":"What are the temperature and humidity readings?"}\n\n# Phan hoi hoi dap da ghi nhan\n{"tra_loi":"The current temperature is 29.9°C, and the humidity is at 83%."}')
H('Truy vấn đối chiếu tập dữ liệu lịch sử')
P('Các truy vấn dùng mốc trước ngày 15/10 để tách tập dữ liệu lịch sử khỏi các bản tin được nhận lại trong lần khởi động kiểm tra. Với phiên đo mới, cần chọn cửa sổ thời gian tương ứng và kiểm tra nguồn gốc bản tin trước khi thống kê.')
code("SELECT topic, COUNT(*) AS so_hang\nFROM cam_bien_data\nWHERE thoi_gian < '2026-10-15'\nGROUP BY topic;\n\nSELECT topic,\n       MIN(CAST(gia_tri AS DECIMAL(8,2))) AS nho_nhat,\n       MAX(CAST(gia_tri AS DECIMAL(8,2))) AS lon_nhat,\n       AVG(CAST(gia_tri AS DECIMAL(8,2))) AS trung_binh\nFROM cam_bien_data\nWHERE thoi_gian < '2026-10-15'\n  AND topic IN ('nha/phongkhach/nhietdo/state',\n                'nha/phongkhach/doam/state')\nGROUP BY topic;")
H('Các lưu ý khi mở rộng cấu hình')
P('Nếu dùng MQTT cùng native API ESPHome nhưng không có native API client kết nối, cần xem xét cấu hình reboot_timeout theo tài liệu MQTT của ESPHome [2]. Khi sửa cực tính ánh sáng, phải kiểm tra cả nhãn frontend và điều kiện điều khiển relay. Khi triển khai từ điện thoại, thay localhost bằng địa chỉ máy chủ phù hợp trong LAN và cấu hình quyền truy cập cần thiết.',
'Để đo độ trễ điều khiển trong lần phát triển tiếp theo, cần đồng bộ thời gian, thêm mã lệnh và lưu riêng các mốc gửi lệnh, nhận lệnh, cập nhật đầu ra, nhận trạng thái. Để đánh giá cảm biến, cần ghi đồng thời số đọc của thiết bị chuẩn. Để đánh giá tiết kiệm điện, cần đo công suất và thời lượng tải ở một cấu hình đối chứng tương ứng.')

# Materialize lists of figures and tables with live page-reference fields.
def populate_list(anchor,items):
    for label,b in items:
        p=D.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.15
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.7),WD_ALIGN_PARAGRAPH.RIGHT)
        p.add_run(label+'\t');field(p,'PAGEREF '+b+' \\h')
        anchor._p.addprevious(p._p)
    anchor._element.getparent().remove(anchor._element)
populate_list(figure_list_anchor,figures)
populate_list(table_list_anchor,tables)
settings=D.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
D.core_properties.title='Thiết kế hệ thống giám sát và điều khiển nhà thông minh tích hợp trợ lý AI'
D.core_properties.author='Nhóm sinh viên D18CNPM5'
D.core_properties.subject='Báo cáo học phần Hệ thống IoT và ứng dụng'
D.core_properties.keywords='ESP32, ESPHome, MQTT, React, Node.js, MySQL, Ollama'
D.save(OUT)
print(OUT)
print('Paragraphs',len(D.paragraphs),'Tables',len(D.tables),'Figures',len(figures))
