from pathlib import Path
import json
from copy import deepcopy
from PIL import Image,ImageOps
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from network_sections import add_network_sections
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path.cwd();WORK=ROOT/'bao-cao/work';IMAGES=WORK/'images'
SOURCE=WORK/'Bao_cao_truoc_khi_bo_sung_giao_dien.docx'
DEST=ROOT/'bao-cao/Bao_cao_IoT_Nha_thong_minh_bo_sung_giao_dien.docx'
D=Document(SOURCE)
body=D._element.body
original=list(body)
start=next(p._p for p in D.paragraphs if p.style.name=='Heading 2' and p.text=='Thiết kế giao diện giám sát')
end=next(p._p for p in D.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHƯƠNG 4 KẾT QUẢ'))
for e in original[original.index(start):original.index(end)]:body.remove(e)
newnodes=[];counter=51000;newfig=7

def move(p):
    el=p._p;body.remove(el);end.addprevious(el);newnodes.append(el);return p
def H(text,level=3):
    return move(D.add_paragraph(text,style=f'Heading {level}'))
def P(*texts):
    for text in texts:
        p=D.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;move(p)
def mark(p,prefix):
    global counter
    counter+=1;name=f'{prefix}_v2_{counter}'
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(counter));a.set(qn('w:name'),name);p._p.insert(0,a)
    a=OxmlElement('w:bookmarkEnd');a.set(qn('w:id'),str(counter));p._p.append(a);return name
def F(file,label,width=15.8,height=17):
    global newfig
    file=IMAGES/file
    if not file.exists():raise FileNotFoundError(f'Required interface image: {file.name}')
    im=Image.open(file);w=min(width,height*im.width/im.height)
    p=D.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    r=p.add_run();r.add_picture(str(file),width=Cm(w));r._r.xpath('.//wp:docPr')[0].set('descr',label);move(p)
    caption=f'Hình 3.{newfig} {label}';newfig+=1
    p=D.add_paragraph(caption,style='Caption');mark(p,'fig');move(p)
def field(p,instruction):
    for typ,txt in [('begin',None),(None,' '+instruction+' '),('separate',None),(None,'1'),('end',None)]:
        r=p.add_run()
        if typ:
            x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),typ);r._r.append(x)
        elif txt.strip()=='1':r.text=txt
        else:x=OxmlElement('w:instrText');x.set(qn('xml:space'),'preserve');x.text=txt;r._r.append(x)

def T(caption,headers,rows,widths):
    p=D.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True;mark(p,'tbl');move(p)
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for column,w in zip(t.columns,widths):column.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'5');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            cp.append(borders)
            if ri==0:
                x=OxmlElement('w:shd');x.set(qn('w:fill'),'E9EFF5');cp.append(x)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.08
                p.paragraph_format.keep_with_next=(ri==0)
                for r in p.runs:r.font.size=Pt(11.5);r.bold=(ri==0)
    body.remove(t._tbl);end.addprevious(t._tbl)

add_network_sections(H,P,F,T)

H('Giao diện web và tổ chức các chức năng',2)
P('Giao diện web là nơi người dùng quan sát và điều khiển toàn bộ mô hình phòng khách. Những giá trị thu được từ cảm biến, trạng thái relay, lịch sử môi trường và thông tin thói quen được đưa về cùng một ứng dụng, giúp người sử dụng theo dõi ngôi nhà mà không phải mở riêng từng công cụ kỹ thuật. Nhóm xây dựng giao diện theo hướng dễ nhận biết trạng thái, ít thao tác và có thể mở rộng khi bổ sung thiết bị.',
'Thanh điều hướng gồm bốn mục: Bảng điều khiển, Thiết bị sử dụng, Thói quen và Trợ lý AI. Cách phân chia này tương ứng với bốn nhu cầu chính: theo dõi tình trạng phòng, tìm hiểu phần cứng, lưu thông tin sinh hoạt và hỏi đáp. Các trang dùng chung ngôn ngữ thiết kế với nền sáng, thẻ nội dung bo góc, biểu tượng trực quan và màu nhấn nhất quán.',
'Các ảnh minh họa dưới đây được chụp từ ứng dụng của nhóm. Dữ liệu lịch sử được sử dụng lại để thể hiện đầy đủ các thành phần của màn hình trong quá trình trình diễn. Phần mô tả đi từ mục đích của mỗi trang đến thao tác của người dùng và kết quả hệ thống trả về.')

H('Bảng điều khiển giám sát tổng hợp')
P('Bảng điều khiển là màn hình đầu tiên khi truy cập website. Bốn thẻ ở phía trên trình bày nhiệt độ, độ ẩm, chuyển động và ánh sáng; mỗi thẻ có biểu tượng, giá trị hoặc nhãn trạng thái và thời gian cập nhật. Người sử dụng có thể nhận biết nhanh đặc điểm môi trường trong phòng trước khi xem sâu hơn phần lịch sử.',
'Ngay bên dưới là khu vực điều khiển relay và trạng thái chế độ thủ công. Biểu đồ và ô hỏi trợ lý được đặt tiếp theo để người dùng vừa theo dõi biến động, vừa đặt câu hỏi khi cần giải thích dữ liệu. Nhãn kết nối trên thanh điều hướng phản ánh kết nối giữa trình duyệt với backend; thời điểm cập nhật trên từng thẻ giúp bổ sung ngữ cảnh khi đang xem lại dữ liệu cũ.')
F('web-dashboard.png','Màn hình Bảng điều khiển của website nhà thông minh')
P('Khi tải trang, frontend lấy dữ liệu gần nhất qua REST API rồi tiếp tục nhận các bản tin WebSocket. Nhờ cơ chế này, các thẻ có thể hiển thị giá trị đã lưu ngay khi người dùng mở website, sau đó được thay đổi khi có thông tin mới. Trong quá trình trình diễn, dữ liệu cũ vẫn giúp làm rõ cách bố trí và mối liên hệ giữa các thành phần của hệ thống.')

H('Biểu đồ theo dõi nhiệt độ và độ ẩm')
P('Biểu đồ chuyển các giá trị rời rạc thành một diễn biến theo thời gian, giúp người dùng quan sát xu hướng thay đổi của môi trường. Đường màu vàng biểu diễn nhiệt độ, còn đường màu xanh biểu diễn độ ẩm. Chú giải ghi rõ đơn vị để phân biệt hai đại lượng khi cùng xuất hiện trong một khung đồ thị.',
'Hai lựa chọn 20 điểm và 50 điểm cho phép thay đổi số mốc gần nhất được hiển thị. Với 20 điểm, giao diện tập trung vào đoạn dữ liệu ngắn để dễ quan sát; với 50 điểm, người dùng có thêm bối cảnh lịch sử. Chức năng này sử dụng số điểm dữ liệu, không tương đương một khoảng thời gian cố định nếu các mẫu được ghi không đều.')
F('web-chart.png','Biểu đồ nhiệt độ và độ ẩm với dữ liệu lịch sử')
P('Khi phân tích một thời điểm trên đồ thị, người dùng đối chiếu trục thời gian, màu của đường và đơn vị trong chú giải. Dữ liệu được ghép từ hai topic độc lập và giới hạn trong 50 mốc gần nhất ở frontend. Đối với các phiên bị gián đoạn, cần xem thời gian của từng mẫu để diễn giải đúng sự thay đổi giữa các điểm.')

H('Điều khiển relay và lựa chọn chế độ hoạt động')
P('Khối relay thể hiện trạng thái đóng cắt và cung cấp một nút chuyển trạng thái. Khi relay đang tắt, nút có nhãn Bật; khi đang bật, nút đổi thành Tắt. Sau thao tác, giao diện chuyển sang trạng thái chờ và cập nhật khi nhận được bản tin relay/state qua WebSocket. Cách phản hồi này giúp người dùng theo dõi tiến trình xử lý ngay tại nơi thao tác.',
'Chế độ thủ công được tách thành một khu vực riêng để cho biết quyền điều khiển đang được ưu tiên. Khi điều khiển LED hoặc còi bằng câu lệnh, backend gửi yêu cầu bật chế độ thủ công trước khi gửi lệnh đầu ra. Nút Quay lại tự động giúp người sử dụng trả quyền xử lý về các quy tắc cảm biến sau khi kết thúc thao tác thủ công.')
F('web-controls.png','Khu vực điều khiển relay và chế độ vận hành')
P('Trạng thái relay trên web tương ứng với công tắc ở GPIO26; lệnh đèn trong chatbot điều khiển LED GPIO25. Việc thể hiện rõ vai trò của hai đầu ra giúp tránh nhầm lẫn khi trình diễn. Trong các lần thử nghiệm của nhóm, thao tác bật tắt được đối chiếu với trạng thái MQTT và quan sát trên mô hình.')

H('Trang giới thiệu thiết bị sử dụng')
P('Trang Thiết bị sử dụng trình bày các linh kiện theo nhóm chức năng, gồm bộ xử lý trung tâm, cảm biến, cơ cấu chấp hành và linh kiện hỗ trợ. Mỗi thẻ có ảnh thật, tên thiết bị và mô tả vai trò trong hệ thống. Đây là cầu nối giữa sản phẩm phần mềm và mô hình phần cứng, giúp người xem hiểu giá trị trên dashboard được tạo ra từ bộ phận nào.',
'Nhóm cảm biến giới thiệu DHT11, PIR HC-SR501 và module LDR. Phần cơ cấu chấp hành trình bày relay, LED và còi. Thông tin trên từng thẻ gắn với cách sử dụng thực tế của dự án, chẳng hạn DHT11 cung cấp hai đại lượng môi trường, PIR tạo sự kiện chuyển động và relay đóng cắt một kênh tải.')
F('web-overview.png','Trang Thiết bị sử dụng: tổng quan dự án và luồng hoạt động')

H('Theo dõi các phiên bản lắp ráp')
P('Bộ lọc phiên bản cho phép người dùng xem phần cứng theo từng giai đoạn phát triển. Phiên bản đầu tập trung vào thu thập cảm biến và relay; phiên bản sau bổ sung LED, còi và chế độ thủ công. Chức năng lọc đi cùng ảnh lắp ráp và ghi chú thay đổi, giúp người đọc theo dõi quá trình hoàn thiện mô hình thay vì chỉ quan sát kết quả cuối cùng.')
P('Trong quá trình thuyết trình, trang này có thể được dùng để chỉ ra vị trí các linh kiện, lý do bổ sung đầu ra và mối liên hệ với các chân GPIO đã thiết kế. Việc lưu ảnh theo phiên bản cũng giúp thuận tiện đối chiếu khi kiểm tra dây nối hoặc tái lắp mô hình.')

H('Trang quản lý thói quen sinh hoạt')
P('Trang Thói quen cho phép lưu các nội dung sinh hoạt theo danh mục. Những mục đã có được hiển thị thành từng nhóm, mỗi nội dung là một nhãn ngắn kèm nút xóa. Khu vực nhập phía dưới gồm tên danh mục, nội dung thói quen và nút thêm, tạo một quy trình thao tác rõ ràng ngay trên cùng màn hình.',
'Người dùng có thể chọn một danh mục đã tồn tại hoặc nhập tên mới, sau đó ghi nội dung cần lưu. Khi thêm thành công, mục vừa tạo được đưa vào nhóm tương ứng; khi xóa một mục, danh sách trên trang được cập nhật. Dữ liệu nằm trong bảng thoi_quen_nguoi_dung nên vẫn có thể sử dụng sau khi đóng và mở lại trình duyệt.')
F('web-habits.png','Màn hình quản lý thói quen theo danh mục')
P('Trong dữ liệu của dự án có 10 mục thuộc 5 danh mục. Các nội dung này được backend đưa vào ngữ cảnh để trợ lý có thể trả lời các câu hỏi liên quan đến thói quen đã ghi. Chức năng hiện tập trung vào ghi nhớ và tham chiếu thông tin; việc tự chuyển một nội dung thành lịch bật tắt thiết bị là hướng mở rộng tiếp theo.')

H('Trang hội thoại với trợ lý AI')
P('Trang Trợ lý AI cung cấp một không gian hội thoại riêng, gồm các tin nhắn người dùng, câu trả lời của trợ lý và vùng nhập câu hỏi. Khi chưa có hội thoại, các câu hỏi gợi ý giúp người dùng bắt đầu với những nội dung hệ thống đang hỗ trợ, như nhiệt độ và độ ẩm gần nhất, gợi ý môi trường, bất thường và thói quen đã lưu.',
'Người dùng nhập câu hỏi rồi chọn Gửi. Trong lúc backend xử lý, giao diện thể hiện trạng thái đang suy nghĩ; khi có kết quả, câu trả lời xuất hiện thành một tin nhắn riêng. Cách trình bày theo lượt hội thoại giúp người dùng dễ đối chiếu câu hỏi với phản hồi và thuận tiện hơn khi trình diễn nhiều tình huống liên tiếp.')
F('web-chatbot.png','Giao diện Trợ lý AI với ô nhập và các câu hỏi gợi ý')
P('Với câu hỏi về nhiệt độ và độ ẩm, backend lấy giá trị gần nhất từ MySQL rồi gửi kèm vào ngữ cảnh cho Ollama. Với các lệnh LED hoặc còi được nhận diện, backend thực hiện nhánh quy tắc và trả lời xác nhận đã gửi lệnh. Như vậy, cùng một giao diện hội thoại hỗ trợ cả khai thác thông tin và điều khiển một số đầu ra, trong khi hai luồng xử lý bên trong vẫn được phân định rõ.',
'Lịch sử tin nhắn trên trang phục vụ việc theo dõi phiên trao đổi ở frontend. Mỗi lần gọi Ollama hiện gửi ngữ cảnh hệ thống và câu hỏi mới; phiên bản này chưa gửi lại đầy đủ các lượt hội thoại trước đó cho mô hình. Khi dữ liệu lịch sử được sử dụng để minh họa, giá trị trong câu trả lời được đối chiếu với tập dữ liệu đã lưu.')

H('Hiển thị gợi ý và cảnh báo trên giao diện')
P('Ngoài các giá trị cảm biến, dashboard còn có vùng thông báo để giải thích những trạng thái cần người dùng chú ý. Khi nhiệt độ hoặc độ ẩm đi qua các ngưỡng đã khai báo, backend tạo nội dung gợi ý, xác định mức thông báo và gửi đến frontend. Màu sắc của thông báo giúp người sử dụng nhận biết mức độ, còn câu mô tả cho biết giá trị và hành động tham khảo.',
'Các thông báo được bổ sung vào đầu danh sách và có nút đóng để người dùng chủ động thu gọn màn hình. Frontend giữ tối đa 5 thông báo gần nhất. Nội dung gợi ý không tự động bật điều hòa, quạt hoặc máy tạo ẩm; các thiết bị đó được nhắc trong câu gợi ý và chưa phải các đầu ra đã cấu hình trong mô hình.')

H('Công cụ cấu hình và kiểm chứng kết nối',2)
P('Bên cạnh website dành cho người dùng, nhóm sử dụng MQTT Explorer và ESPHome Device Builder để thiết lập, quan sát và đối chiếu hoạt động bên trong hệ thống. Website cho thấy kết quả ở tầng ứng dụng; MQTT Explorer cho thấy topic và payload trên broker; ESPHome Device Builder thể hiện cấu hình của nút ESP32. Ba góc nhìn bổ sung cho nhau trong quá trình lắp ráp và thử nghiệm.')

H('Thiết lập kết nối bằng MQTT Explorer')
P('Ảnh cấu hình của nhóm cho thấy MQTT Explorer kết nối tới broker trên máy cục bộ với giao thức mqtt://, địa chỉ 127.0.0.1 và cổng 1883. Sau khi kết nối, công cụ hiển thị các nhánh topic mà broker chuyển đến. Đây là bước kiểm tra trực tiếp đường truyền trước khi đối chiếu dữ liệu trên dashboard.',
'Trong quy trình thử nghiệm, nhóm kiểm tra lần lượt địa chỉ broker, cổng dịch vụ và topic đang được công bố. Khi dashboard chưa thay đổi như mong muốn, việc quan sát MQTT Explorer giúp xác định bản tin có xuất hiện trên broker hay không, qua đó thu hẹp phạm vi kiểm tra giữa thiết bị, máy chủ và giao diện.')
F('mqtt-connection-user.png','Cấu hình kết nối broker cục bộ trên MQTT Explorer',height=15)

H('Quan sát cây topic và các giá trị cảm biến')
P('Nhánh nha/phongkhach tổ chức dữ liệu theo từng loại: nhietdo, doam, chuyendong, anhsang và relay. Trong ảnh do nhóm ghi lại, nhiệt độ hiển thị 32,5°C, độ ẩm 59%, chuyển động OFF, ánh sáng ON và relay OFF. Mỗi giá trị nằm trong một topic trạng thái riêng, phù hợp với giao ước mà backend sử dụng.',
'Cây topic giúp kiểm tra nhanh cả cấu trúc đặt tên lẫn nội dung payload. Khi nhiệt độ hoặc độ ẩm thay đổi ở thiết bị, người kiểm thử có thể quan sát bản tin tại đây rồi đối chiếu giá trị tương ứng trên website. Ảnh thể hiện một thời điểm trong quá trình triển khai, vì vậy các giá trị minh họa có thể khác với bộ dữ liệu được nạp lại khi chụp các màn hình web.')
F('mqtt-topics-focused.png','Nhánh topic phòng khách và dữ liệu ghi nhận trên MQTT Explorer',width=13.8,height=15)

H('Xây dựng cấu hình bằng ESPHome Device Builder').paragraph_format.page_break_before=True
P('ESPHome Device Builder cho phép quản lý các thành phần của nút cam-bien-phong-khach theo nhóm như ESPHome Core, ESP32 Platform, Logger, MQTT Client, Native API, OTA và WiFi. Vùng cấu hình ở bên phải hiển thị YAML tương ứng, giúp nhóm đối chiếu các tham số của thiết bị trong cùng một không gian làm việc.',
'Ảnh chụp thể hiện tên nút, biến thể ESP32, framework esp-idf và địa chỉ broker. Các trường mật khẩu và khóa được che trên giao diện. Từ cấu hình này, nhóm bổ sung cảm biến, công tắc và các hành động tự động đã trình bày trong các phần trước. Cấu hình đầy đủ đã tách thông tin xác thực được đặt ở Phụ lục A để thuận tiện tham khảo.')
F('esphome-builder-user.png','Giao diện ESPHome Device Builder dùng để cấu hình nút ESP32')
P('Việc kết hợp ảnh mô hình, các màn hình website và ảnh công cụ kỹ thuật giúp mô tả sản phẩm ở cả ba mức: thiết bị vật lý, truyền thông và tương tác người dùng. Qua đó, quá trình triển khai được thể hiện bằng các thành phần có thể quan sát, thay vì chỉ dừng ở mô tả kiến trúc hoặc danh sách chức năng.')

# Rewrite selected transitions to make the report read as a course report in the group's voice.
replacements={
'Trong một phòng sinh hoạt, người sử dụng thường cần biết nhiệt độ, độ ẩm và tình trạng thiết bị, đồng thời muốn bật tắt các cơ cấu chấp hành mà không phải thao tác trực tiếp tại từng công tắc. Nếu dữ liệu chỉ được quan sát tại chỗ, việc theo dõi diễn biến theo thời gian và đối chiếu các sự kiện trở nên khó khăn. Đây là bài toán phù hợp để vận dụng kiến thức IoT từ thu thập tín hiệu đến truyền dữ liệu, lưu trữ và xây dựng giao diện.':
'Một hệ thống nhà thông minh cần giúp người sử dụng trả lời hai câu hỏi thiết thực: không gian trong nhà đang ở trạng thái nào và có thể tác động đến thiết bị bằng cách nào. Với một phòng khách, nhu cầu đó được thể hiện qua việc theo dõi nhiệt độ, độ ẩm, sự xuất hiện của chuyển động, trạng thái ánh sáng và khả năng bật tắt các đầu ra. Đề tài được xây dựng từ những nhu cầu cụ thể này để kết nối kiến thức IoT với một sản phẩm có thể quan sát và thao tác trực tiếp.',
'Các thử nghiệm chức năng trước đây được nhóm xác nhận thành công, bao gồm giám sát cảm biến, điều khiển tự động và thủ công, lưu lịch sử, cập nhật giao diện và hỏi đáp AI. Lần kiểm tra ngày 15/10/2026 tiếp tục xác nhận kết nối MQTT, REST API, WebSocket, quá trình build giao diện và phản hồi chatbot. Báo cáo đồng thời phân tích giới hạn của kiến trúc hiện tại và đề xuất hướng hoàn thiện.':
'Nhóm đã thử nghiệm thành công các chức năng giám sát cảm biến, điều khiển tự động và thủ công, lưu lịch sử, cập nhật giao diện và hỏi đáp AI. Các màn hình web, ảnh mô hình và dữ liệu MQTT minh họa cho từng kết quả triển khai. Lần kiểm tra ngày 15/10/2026 tiếp tục xác nhận khả năng kết nối, truy vấn và phản hồi của các dịch vụ phần mềm. Từ kết quả đạt được, báo cáo phân tích những giới hạn còn lại và đề xuất các bước phát triển tiếp theo.',
'Trong các lần chạy trước, chatbot đã trả lời tốt theo xác nhận của nhóm. Khi kiểm tra lại, câu hỏi “What are the temperature and humidity readings? Answer in one sentence.” nhận phản hồi HTTP 200 với nội dung “The current temperature is 29.9°C, and the humidity is at 83%.” Hai giá trị trùng với các giá trị trạng thái đã được nạp vào backend.':
'Chatbot đã được nhóm sử dụng thành công để hỏi về dữ liệu cảm biến và nội dung tham chiếu của hệ thống. Trong lần kiểm tra lại, câu hỏi “What are the temperature and humidity readings? Answer in one sentence.” nhận phản hồi HTTP 200 với nội dung “The current temperature is 29.9°C, and the humidity is at 83%.” Hai giá trị trùng với trạng thái đã nạp vào backend, thể hiện được luồng lấy dữ liệu, tạo ngữ cảnh và sinh câu trả lời. Màn hình hội thoại ở Chương 3 minh họa cách người dùng tiếp cận chức năng này trên website.'
}
for p in D.paragraphs:
    if p.text in replacements:
        p.text=replacements[p.text]

# Add primary technical references for the expanded network discussion.
ref_end=next(p for p in D.paragraphs if p.style.name=='Heading 1' and p.text.startswith('PHỤ LỤC A CẤU HÌNH'))
for text in [
    '[10] OASIS. MQTT Version 3.1.1, OASIS Standard, 29/10/2014. Các mục QoS, retained message và topic filter. Truy cập ngày 15/10/2026. https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/os/mqtt-v3.1.1-os.html',
    '[11] Docker. Port publishing and mapping. Tài liệu trực tuyến. Truy cập ngày 15/10/2026. https://docs.docker.com/engine/network/port-publishing/',
    '[12] Vite. Server Options. Tài liệu trực tuyến. Truy cập ngày 15/10/2026. https://vite.dev/config/server-options']:
    p=D.add_paragraph(text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.15
    for r in p.runs:r.font.size=Pt(12)
    body.remove(p._p);ref_end._p.addprevious(p._p)

# Rebuild the figure directory from real captions and bookmark page references.
paras=D.paragraphs
fstart=next(p for p in paras if p.text=='DANH MỤC HÌNH ẢNH')
fend=next(p for p in paras[paras.index(fstart)+1:] if p.text=='DANH MỤC BẢNG BIỂU')
els=list(body);a=els.index(fstart._p);b=els.index(fend._p)
# Preserve the page break that introduces the table directory.
for el in els[a+1:b]:
    if not el.xpath('.//w:br[@w:type="page"]'):body.remove(el)
for p in D.paragraphs:
    if p.style.name=='Caption' and p.text.startswith('Hình '):
        bookmarks=p._p.xpath('.//w:bookmarkStart')
        if bookmarks:name=bookmarks[0].get(qn('w:name'))
        else:name=mark(p,'fig')
        row=D.add_paragraph();row.paragraph_format.first_line_indent=Cm(0);row.paragraph_format.space_after=Pt(6);row.paragraph_format.line_spacing=1.15
        row.paragraph_format.tab_stops.add_tab_stop(Cm(15.7),WD_TAB_ALIGNMENT.RIGHT)
        row.add_run(p.text+'\t');field(row,'PAGEREF '+name+' \\h')
        body.remove(row._p);fend._p.addprevious(row._p)
# Move the directory page break after the newly generated rows.
els=list(body);a=els.index(fstart._p);b=els.index(fend._p)
breaks=[el for el in els[a+1:b] if el.xpath('.//w:br[@w:type="page"]')]
for el in breaks:fend._p.addprevious(el)

# Rebuild the table directory as well, including the three deployment tables.
paras=D.paragraphs
tstart=next(p for p in paras if p.text=='DANH MỤC BẢNG BIỂU')
tend=next(p for p in paras[paras.index(tstart)+1:] if p._p.xpath('./w:pPr/w:sectPr'))
els=list(body);a=els.index(tstart._p);b=els.index(tend._p)
for el in els[a+1:b]:body.remove(el)
for p in D.paragraphs:
    if p.style.name=='Caption' and p.text.startswith('Bảng '):
        bookmarks=p._p.xpath('.//w:bookmarkStart');name=bookmarks[0].get(qn('w:name')) if bookmarks else mark(p,'tbl')
        row=D.add_paragraph();row.paragraph_format.first_line_indent=Cm(0);row.paragraph_format.space_after=Pt(6);row.paragraph_format.line_spacing=1.15
        row.paragraph_format.tab_stops.add_tab_stop(Cm(15.7),WD_TAB_ALIGNMENT.RIGHT)
        row.add_run(p.text+'\t');field(row,'PAGEREF '+name+' \\h')
        body.remove(row._p);tend._p.addprevious(row._p)
D.core_properties.title='Thiết kế hệ thống giám sát và điều khiển nhà thông minh tích hợp trợ lý AI'
D.core_properties.comments=''
D.save(DEST)
print(DEST)
print('Images',len(D.inline_shapes),'Figures',sum(p.style.name=='Caption' and p.text.startswith('Hình ') for p in D.paragraphs))
